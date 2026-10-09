import openpyxl
import builtins
import os
import queue
import re
import time
import threading
import traceback
import unicodedata
from pathlib import Path
from datetime import datetime, timedelta
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.common.keys import Keys


PASTA = Path(__file__).resolve().parent
PASTA_ERROS = PASTA / "erros"
ARQUIVO = PASTA / "PlanosdeTratamentoVitória.xlsx"
ABA = "Plano de Tratamento migrar"
CLINICA = "GRANDE VITORIA - GRUPO MANDIC"
URL_BASE = "https://app.clinicanasnuvens.com.br"

COL_STATUS = "Status"
STATUS_FINALIZADOS = ("cadastrado", "erro")

LIMITE_NOME = 30
LIMITE_DESCRICAO = 150
LIMITE_QUANTIDADE = 100

# Profissional do plano: tenta o Aluno, depois o Coordenador e, se nenhum der certo,
# usa o gerente da unidade abaixo
PROFISSIONAL = "ENZO GABRIEL MENDES PARANAGUA"
COL_ALUNO = "Aluno Associado"
COL_COORDENADOR = "Coordenador"

# True = só imprime a estrutura de cada paciente, sem cadastrar nem gravar Status
MODO_REVISAO = False

# Quantos pacientes processar nesta execução (None = todos). Use 1 para testar.
LIMITE_PACIENTES = None

# Quantos navegadores trabalham ao mesmo tempo (cada Chrome usa ~0,5–1 GB de RAM)
NUM_NAVEGADORES = 1
# Segundos entre o login de um navegador e o do próximo
INTERVALO_ENTRE_LOGINS = 15

# Só um navegador por vez lê/grava a planilha; e as linhas do terminal não se misturam
trava_planilha = threading.Lock()
trava_print = threading.Lock()
trava_progresso = threading.Lock()


def print(*args, sep=" ", end="\n", **kwargs):
    """print com o nome do navegador ([N1], [N2]...) no começo de cada linha."""
    nome = threading.current_thread().name
    mensagem = sep.join(str(a) for a in args)
    if re.fullmatch(r"N\d+", nome):
        mensagem = "\n".join(f"[{nome}] {linha}" if linha else linha for linha in mensagem.split("\n"))
    with trava_print:
        builtins.print(mensagem, end=end, **kwargs)


def texto(valor):
    return str(valor).strip() if valor is not None else ""


def normalizar(valor):
    """Minúsculas, sem acentos e com espaços simples — para comparar nomes."""
    valor = unicodedata.normalize("NFKD", texto(valor))
    valor = "".join(c for c in valor if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", valor).strip().lower()


def carregar_pacientes():
    """Lê a aba 'Plano de Tratamento' e monta a estrutura:

    pacientes[PG] = {
        "nomePaciente": str,
        "linhas": [nº das linhas na planilha],
        "planos": {
            <Perfil Tratamento>: {
                "linhas": [nº das linhas do plano],
                "orcamentos": [orçamentos de origem],
                "aluno": str, "coordenador": str,  # 1º valor preenchido do plano
                "procedimentos": {
                    <Procedimento>: {"especialidade": str, "quantidade": int},
                },
            },
        },
    }

    Cada Perfil Tratamento vira um plano de tratamento. Se o mesmo procedimento
    aparece mais de uma vez no plano (ex.: orçamentos diferentes com o mesmo
    perfil), as quantidades são somadas. Linhas com Status 'cadastrado' ou
    'erro' são puladas (o Status é gravado por plano).
    """
    wb = openpyxl.load_workbook(ARQUIVO, read_only=True)
    linhas = wb[ABA].iter_rows(values_only=True)
    cabecalho = [texto(c) for c in next(linhas)]
    idx = {nome: i for i, nome in enumerate(cabecalho)}

    pacientes = {}

    for num_linha, linha in enumerate(linhas, start=2):
        pg = texto(linha[idx["PG"]])
        if not pg:
            continue

        if COL_STATUS in idx and texto(linha[idx[COL_STATUS]]).lower() in STATUS_FINALIZADOS:
            continue

        perfil = texto(linha[idx["Perfil Tratamento"]])
        procedimento = texto(linha[idx["Procedimento (sistema novo)"]])
        especialidade = texto(linha[idx["Especialidade"]])
        orcamento = texto(linha[idx["Orçamento"]])
        quantidade = int(linha[idx["Quantidade"]] or 0)

        if pg not in pacientes:
            pacientes[pg] = {
                "nomePaciente": texto(linha[idx["Paciente"]]),
                "linhas": [],
                "planos": {},
            }
        pacientes[pg]["linhas"].append(num_linha)

        plano = pacientes[pg]["planos"].setdefault(
            perfil, {"linhas": [], "orcamentos": [], "aluno": "", "coordenador": "", "procedimentos": {}}
        )
        plano["linhas"].append(num_linha)
        if orcamento not in plano["orcamentos"]:
            plano["orcamentos"].append(orcamento)

        for chave, coluna in (("aluno", COL_ALUNO), ("coordenador", COL_COORDENADOR)):
            if not plano[chave] and coluna in idx:
                plano[chave] = texto(linha[idx[coluna]])

        proc = plano["procedimentos"].setdefault(procedimento, {"especialidade": especialidade, "quantidade": 0})
        proc["quantidade"] += quantidade

    wb.close()
    return pacientes


def marcar_status(linhas, valor):
    """Grava `valor` na coluna Status das linhas informadas — um navegador por vez."""
    with trava_planilha:
        _marcar_status(linhas, valor)


def _marcar_status(linhas, valor):
    wb = openpyxl.load_workbook(ARQUIVO)
    ws = wb[ABA]

    cabecalho = [texto(c.value) for c in ws[1]]
    if COL_STATUS in cabecalho:
        col_status = cabecalho.index(COL_STATUS) + 1
    else:
        col_status = len(cabecalho) + 1
        ws.cell(row=1, column=col_status, value=COL_STATUS)

    for linha in linhas:
        ws.cell(row=linha, column=col_status, value=valor)

    # Salva numa cópia e só depois troca pelo original: se o script for
    # interrompido no meio do salvamento, a planilha original fica intacta.
    temporario = ARQUIVO.with_name(f"~tmp_{ARQUIVO.name}")
    wb.save(temporario)
    os.replace(temporario, ARQUIVO)


def titulo_plano(perfil, plano):
    return f"Plano '{perfil}' (orçamentos: {', '.join(plano['orcamentos'])})"


def imprimir_paciente(pg, dados):
    print(f"  Paciente: {dados['nomePaciente']} (PG: {pg}) | linhas na planilha: {dados['linhas']}")
    for perfil, plano in dados["planos"].items():
        print(f"  - {titulo_plano(perfil, plano)}")
        print(f"      Aluno: {plano['aluno'] or '-'} | Coordenador: {plano['coordenador'] or '-'}")
        for procedimento, info in plano["procedimentos"].items():
            print(f"      {info['quantidade']}x {procedimento} [{info['especialidade']}]")


def imprimir_amostra(pacientes, quantidade=10):
    for i, (pg, dados) in enumerate(pacientes.items()):
        if i >= quantidade:
            break

        print()
        imprimir_paciente(pg, dados)


def tamanho_textarea(valor):
    # Por segurança, conta a quebra de linha como 2 caracteres (\r\n) no maxlength
    return len(valor) + valor.count("\n")


def montar_descricao(perfil, plano):
    """Título do plano + procedimentos, cabendo em LIMITE_DESCRICAO caracteres.

    Tenta com a especialidade; se não couber, sem ela; se ainda não couber, corta com '…'.
    """
    for com_especialidade in (True, False):
        linhas = [titulo_plano(perfil, plano)]
        for procedimento, info in plano["procedimentos"].items():
            linha = f"{info['quantidade']}x {procedimento}"
            if com_especialidade:
                linha += f" [{info['especialidade']}]"
            linhas.append(linha)

        descricao = "\n".join(linhas)
        if tamanho_textarea(descricao) <= LIMITE_DESCRICAO:
            return descricao

    while tamanho_textarea(descricao + "…") > LIMITE_DESCRICAO:
        descricao = descricao[:-1]
    return descricao.rstrip() + "…"


def formatar_duracao(segundos):
    segundos = int(segundos)
    horas, resto = divmod(segundos, 3600)
    minutos, segundos = divmod(resto, 60)
    return f"{horas:02d}:{minutos:02d}:{segundos:02d}"


def relogio_previsao(progresso, parar, intervalo=60):
    """A cada `intervalo` segundos, imprime o tempo restante estimado para terminar."""
    while not parar.wait(intervalo):
        concluidos = progresso["concluidos"]
        total = progresso["total"]
        decorrido = time.time() - progresso["inicio"]

        if concluidos == 0:
            print(f"\n⏱  [{formatar_duracao(decorrido)} decorrido] Calculando previsão...")
            continue

        media = decorrido / concluidos
        restante = media * (total - concluidos)
        termino = datetime.now() + timedelta(seconds=restante)
        print(
            f"\n⏱  {concluidos}/{total} pacientes | média {media:.1f}s/paciente | "
            f"restante: {formatar_duracao(restante)} | término previsto: {termino:%d/%m %H:%M}"
        )


def salvar_print_erro(driver, nome):
    try:
        PASTA_ERROS.mkdir(exist_ok=True)
        nome = re.sub(r"[^\w\-]+", "_", nome)
        driver.save_screenshot(str(PASTA_ERROS / f"{nome}.png"))
        print(f"  Screenshot salvo em {PASTA_ERROS / (nome + '.png')}")
    except Exception:
        pass


# ── Passos na plataforma ───────────────────────────────────────────────────────

def clicar(driver, elemento):
    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", elemento)
    driver.execute_script("arguments[0].click();", elemento)


def clicar_mouse(driver, elemento):
    """Clique 'real' (mousedown/mouseup/click) — necessário para select2 e autocomplete."""
    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", elemento)
    ActionChains(driver).move_to_element(elemento).pause(0.2).click(elemento).perform()


def localizar_paciente(driver, wait, pg):
    """Passo 1: filtra a lista de pacientes pelo PG e devolve a URL do perfil do paciente."""
    driver.get(f"{URL_BASE}/pacientes")
    time.sleep(2)

    clicar(driver, wait.until(EC.element_to_be_clickable((By.ID, "btn-drawer-filtrosearch-paciente"))))
    time.sleep(1)

    campo_pg = wait.until(EC.element_to_be_clickable((By.ID, "pesquisa-numero-controle")))
    campo_pg.clear()
    campo_pg.send_keys(pg)

    # A lista já vem preenchida: guarda o 1º resultado atual para esperar ele ser substituído
    seletor_link = (By.CSS_SELECTOR, "a[aria-label='Visualizar perfil do paciente']")
    antigos = driver.find_elements(*seletor_link)

    clicar(driver, wait.until(EC.element_to_be_clickable((By.ID, "btn-visualizar-pacientes"))))

    if antigos:
        try:
            wait.until(EC.staleness_of(antigos[0]))
        except TimeoutException:
            pass
    time.sleep(2)

    try:
        wait.until(EC.presence_of_element_located(seletor_link))
    except TimeoutException:
        raise Exception(f"paciente com PG {pg} não encontrado")

    links = driver.find_elements(*seletor_link)
    if len(links) > 1:
        print(f"  Atenção: {len(links)} pacientes com o PG {pg} — usando o primeiro.")

    return links[0].get_attribute("href")


def abrir_novo_plano(driver, wait, url_paciente):
    """Passo 2: abre o paciente > Comercial > Planos de tratamento > Adicionar plano."""
    driver.get(url_paciente)
    time.sleep(2)

    comercial = wait.until(EC.presence_of_element_located(
        (By.XPATH, "//span[contains(@class,'menu-title') and normalize-space()='Comercial']")
    ))
    clicar(driver, comercial)
    time.sleep(1)

    item_planos = wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, "div.menu-link.planos-tratamento")))
    clicar(driver, item_planos)
    time.sleep(2)

    botao_novo = wait.until(EC.presence_of_element_located(
        (By.CSS_SELECTOR, "a[href*='/plano-tratamento/novo/']")
    ))
    clicar(driver, botao_novo)

    wait.until(EC.element_to_be_clickable((By.ID, "pti-nome")))
    time.sleep(1)


def selecionar_profissional(driver, wait, plano):
    """Tenta o Aluno, depois o Coordenador e, por último, o PROFISSIONAL fixo (gerente)."""
    campo = wait.until(EC.element_to_be_clickable((By.ID, "pbv-profissional")))

    candidatos = [("Aluno", plano["aluno"]), ("Coordenador", plano["coordenador"])]
    for papel, nome in candidatos:
        if not nome:
            print(f"    {papel}: vazio na planilha — pulando.")
            continue
        print(f"    Tentando {papel}: {nome}")
        try:
            valor = selecionar_autocomplete(driver, campo, nome, exigir_nome=True, tentativas=2)
            print(f"    Profissional ({papel}): {valor}")
            return
        except Exception as e:
            print(f"    {papel} '{nome}' não selecionado ({resumo_erro(e)}) — tentando o próximo.")

    valor = selecionar_autocomplete(driver, campo, PROFISSIONAL, preferir_exato=False)
    print(f"    Profissional (gerente da unidade): {valor}")


def preencher_cabecalho(driver, wait, nome, descricao, plano):
    """Passo 3: nome (máx. 30), descrição (máx. 150) e profissional do plano."""
    campo_nome = wait.until(EC.element_to_be_clickable((By.ID, "pti-nome")))
    campo_nome.clear()
    campo_nome.send_keys(nome[:LIMITE_NOME])

    campo_descricao = driver.find_element(By.ID, "pti-descricao")
    campo_descricao.clear()
    campo_descricao.send_keys(descricao)

    tamanho = len(campo_descricao.get_attribute("value") or "")
    if tamanho > LIMITE_DESCRICAO:
        raise Exception(f"descrição com {tamanho} caracteres (limite {LIMITE_DESCRICAO})")

    selecionar_profissional(driver, wait, plano)


def campo_procedimento_vazio(driver):
    campos = [
        c for c in driver.find_elements(By.CSS_SELECTOR, "input.tipo-procedimento")
        if c.is_displayed() and not c.get_attribute("value")
    ]
    return campos[-1] if campos else False


def container_do_campo(campo):
    """Caixa (modal) onde está o campo do procedimento; se não for modal, o form."""
    for xpath in ("./ancestor::div[contains(concat(' ', @class, ' '), ' modal ')][1]", "./ancestor::form[1]"):
        encontrados = campo.find_elements(By.XPATH, xpath)
        if encontrados:
            return encontrados[0]
    raise Exception("caixa do procedimento não encontrada")


def itens_autocomplete(driver):
    return [
        li for li in driver.find_elements(By.CSS_SELECTOR, "ul.ui-autocomplete li.ui-menu-item")
        if li.is_displayed()
    ]


def resumo_erro(e):
    """Tipo + 1ª linha da mensagem (as do Selenium trazem um stacktrace enorme)."""
    mensagem = texto(getattr(e, "msg", None) or e).split("\n")[0]
    return f"{type(e).__name__}: {mensagem}" if mensagem else type(e).__name__


def digitar_no_campo(driver, campo, valor):
    """Foca, limpa e digita no campo; se o clique for bloqueado (algo por cima), foca via JS."""
    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", campo)
    try:
        campo.click()
    except Exception:
        driver.execute_script("arguments[0].focus();", campo)
    campo.clear()
    campo.send_keys(valor)


def selecionar_autocomplete(driver, campo, nome, preferir_exato=True, tentativas=3, exigir_nome=False):
    """Digita `nome` num campo autocomplete e escolhe a opção.

    Com `preferir_exato`, prefere a opção com o nome exato; senão pega a primeira.
    Com `exigir_nome`, só aceita opção que contenha o nome digitado (senão falha,
    em vez de pegar outra pessoa qualquer).
    Tenta primeiro com clique do mouse e, se não 'grudar', pelo teclado
    (seta para baixo + Enter).
    """
    for tentativa in range(1, tentativas + 1):
        try:
            digitar_no_campo(driver, campo, nome)
        except Exception as e:
            print(f"    Tentativa {tentativa}: não consegui digitar '{nome}' no campo ({resumo_erro(e)}).")
            time.sleep(2)
            continue

        try:
            WebDriverWait(driver, 20).until(lambda d: itens_autocomplete(d) or False)
        except TimeoutException:
            print(f"    Tentativa {tentativa}: nenhuma opção apareceu para '{nome}'.")
            continue
        time.sleep(1)  # espera a lista terminar de carregar

        itens = itens_autocomplete(driver)
        if not itens:
            continue

        textos = [texto(li.text).split("\n")[0] for li in itens]
        posicao = 0
        if exigir_nome:
            compativeis = [i for i, t in enumerate(textos) if normalizar(nome) in normalizar(t)]
            if not compativeis:
                raise Exception(f"nenhuma opção corresponde a '{nome}' (opções: {textos[:5]})")
            exatos = [i for i in compativeis if normalizar(textos[i]) == normalizar(nome)]
            posicao = (exatos or compativeis)[0]
        elif preferir_exato:
            posicao = next((i for i, t in enumerate(textos) if normalizar(t) == normalizar(nome)), 0)
        escolhido = textos[posicao]

        try:
            if tentativa == 1:
                clicar_mouse(driver, itens[posicao])
            else:
                for _ in range(posicao + 1):
                    campo.send_keys(Keys.ARROW_DOWN)
                    time.sleep(0.2)
                campo.send_keys(Keys.RETURN)
        except Exception as e:
            print(f"    Tentativa {tentativa}: falha ao clicar na opção ({e}).")
            continue
        time.sleep(1)

        valor = texto(campo.get_attribute("value"))
        if valor and not itens_autocomplete(driver):
            if normalizar(valor) != normalizar(nome):
                print(f"    Atenção: selecionado '{valor}' para '{nome}'.")
            return valor

        print(f"    Tentativa {tentativa}: opção '{escolhido}' não ficou selecionada.")

    raise Exception(f"não foi possível selecionar '{nome}'")


def preencher_quantidade(driver, caixa, quantidade):
    """Preenche o campo de quantidade da caixa do procedimento (máx. LIMITE_QUANTIDADE)."""
    if not 1 <= quantidade <= LIMITE_QUANTIDADE:
        raise Exception(f"quantidade {quantidade} fora do limite (1 a {LIMITE_QUANTIDADE})")

    campo = WebDriverWait(driver, 10).until(lambda d: next(
        (c for c in caixa.find_elements(By.CSS_SELECTOR, "input.quantidade-atendimentos") if c.is_displayed()),
        False,
    ))

    campo.click()
    driver.execute_script("arguments[0].select();", campo)
    campo.send_keys(Keys.BACKSPACE)
    campo.send_keys(str(quantidade))
    time.sleep(0.3)

    # Plano B: grava direto no campo e dispara os eventos
    if texto(campo.get_attribute("value")) != str(quantidade):
        driver.execute_script("""
            arguments[0].value = arguments[1];
            arguments[0].dispatchEvent(new Event('input', {bubbles: true}));
            arguments[0].dispatchEvent(new Event('change', {bubbles: true}));
        """, campo, str(quantidade))
        time.sleep(0.3)

    if texto(campo.get_attribute("value")) != str(quantidade):
        raise Exception(f"não foi possível preencher a quantidade {quantidade} (campo ficou '{campo.get_attribute('value')}')")


def selecionar_regiao(driver, caixa):
    """Troca 'Individual (por dente)' por 'Região' no select2 da caixa."""
    selecao = WebDriverWait(driver, 10).until(lambda d: next(
        (s for s in caixa.find_elements(By.CSS_SELECTOR, ".select2-selection.tipo-procedimento-odonto") if s.is_displayed()),
        False,
    ))
    clicar_mouse(driver, selecao)

    try:
        opcao = WebDriverWait(driver, 5).until(lambda d: next(
            (o for o in d.find_elements(By.CSS_SELECTOR, ".select2-container--open .select2-results__option")
             if normalizar(o.text).startswith("regiao")),
            False,
        ))
        clicar_mouse(driver, opcao)
        time.sleep(0.5)
    except TimeoutException:
        pass

    if "regiao" in normalizar(selecao.text):
        return

    # Plano B: seleciona direto no <select> escondido e dispara o change
    ok = driver.execute_script("""
        var select = arguments[0].querySelector('select.tipo-procedimento-odonto');
        if (!select) return false;
        var normalizar = function (t) {
            return (t || '').normalize('NFD').replace(/[\\u0300-\\u036f]/g, '').trim().toLowerCase();
        };
        for (var i = 0; i < select.options.length; i++) {
            var opcao = select.options[i];
            if (normalizar(opcao.text).indexOf('regiao') === 0 || normalizar(opcao.value).indexOf('regiao') === 0) {
                if (window.jQuery) { jQuery(select).val(opcao.value).trigger('change'); }
                else { select.value = opcao.value; select.dispatchEvent(new Event('change', {bubbles: true})); }
                return true;
            }
        }
        return false;
    """, caixa)
    time.sleep(0.5)

    if not ok:
        raise Exception("opção 'Região' não encontrada")


def selecionar_arcada(driver, caixa):
    radio = WebDriverWait(driver, 10).until(lambda d: next(
        (r for r in caixa.find_elements(By.CSS_SELECTOR, "input[type='radio'][value='ARCADA_SUPERIOR_E_INFERIOR']")
         if r.find_element(By.XPATH, "./ancestor::label[1]").is_displayed()),
        False,
    ))
    clicar(driver, radio)
    time.sleep(0.3)

    if not radio.is_selected():
        raise Exception("não foi possível marcar 'Arcada superior e inferior'")


def salvar_caixa(driver, caixa, campo):
    botoes = [
        b for b in caixa.find_elements(By.CSS_SELECTOR, "button[type='submit']")
        if b.is_displayed() and b.get_attribute("id") != "salvar-novo-plano" and "Salvar" in b.text
    ]
    if not botoes:
        raise Exception("botão Salvar do procedimento não encontrado")

    clicar(driver, botoes[-1])

    # A caixa fecha quando o procedimento é salvo
    try:
        WebDriverWait(driver, 10).until(lambda d: not campo.is_displayed())
    except TimeoutException:
        raise Exception("a caixa do procedimento não fechou após salvar (erro de validação?)")
    except Exception:
        pass  # campo removido da página = salvo
    time.sleep(0.5)


def adicionar_procedimento(driver, wait, nome, quantidade):
    """Adicionar > Adicionar procedimento > busca > quantidade > Região > Arcada sup. e inf. > Salvar."""
    botao_adicionar = wait.until(EC.element_to_be_clickable(
        (By.XPATH, "//button[@data-kt-menu-trigger='click' and contains(normalize-space(.),'Adicionar')]")
    ))
    clicar_mouse(driver, botao_adicionar)

    # Espera o menu abrir; se o clique do mouse não abriu, tenta via JS
    try:
        opcao = WebDriverWait(driver, 3).until(EC.visibility_of_element_located((By.ID, "adicionar-procedimento-geral")))
    except TimeoutException:
        clicar(driver, botao_adicionar)
        time.sleep(0.5)
        opcao = driver.find_element(By.ID, "adicionar-procedimento-geral")

    clicar(driver, opcao)

    campo = WebDriverWait(driver, 10).until(campo_procedimento_vazio)
    caixa = container_do_campo(campo)
    time.sleep(0.5)

    selecionar_autocomplete(driver, campo, nome)
    preencher_quantidade(driver, caixa, quantidade)
    selecionar_regiao(driver, caixa)
    selecionar_arcada(driver, caixa)
    salvar_caixa(driver, caixa, campo)


def salvar_plano(driver, wait):
    url_antes = driver.current_url
    clicar(driver, wait.until(EC.element_to_be_clickable((By.ID, "salvar-novo-plano"))))

    try:
        WebDriverWait(driver, 20).until(lambda d: d.current_url != url_antes)
    except TimeoutException:
        raise Exception("o plano não foi salvo (a página não saiu do cadastro)")
    time.sleep(1)


def cadastrar_plano(driver, wait, url_paciente, perfil, plano):
    abrir_novo_plano(driver, wait, url_paciente)
    preencher_cabecalho(driver, wait, perfil, montar_descricao(perfil, plano), plano)

    for procedimento, info in plano["procedimentos"].items():
        print(f"    + {info['quantidade']}x {procedimento}")
        adicionar_procedimento(driver, wait, procedimento, info["quantidade"])

    salvar_plano(driver, wait)


def abrir_sessao():
    """Abre um Chrome, faz login e seleciona a clínica. Devolve (driver, wait)."""
    opcoes = webdriver.ChromeOptions()
    # opcoes.add_argument("--headless=new")
    driver = webdriver.Chrome(options=opcoes)
    driver.maximize_window()
    wait = WebDriverWait(driver, 15)

    try:
        # ── Login ──────────────────────────────────────────────────────────────
        driver.get(f"{URL_BASE}/")

        campo_email = wait.until(EC.presence_of_element_located(
            (By.XPATH, "/html/body/div[1]/div/div[2]/div[1]/div/div[3]/form/div[1]/input")
        ))
        campo_email.send_keys("mmerlo++++dpegoraro@bionexo.com")
        driver.find_element(By.XPATH, "/html/body/div[1]/div/div[2]/div[1]/div/div[3]/form/div[2]/button").click()

        campo_senha = wait.until(EC.element_to_be_clickable((By.ID, "password")))
        campo_senha.click()
        campo_senha.send_keys("CNN@foguete")
        driver.find_element(By.XPATH, "/html/body/div[1]/div/div[2]/div[1]/div/div[4]/form/div[4]/button").click()
        time.sleep(5)

        # ── Seleciona a clínica ────────────────────────────────────────────────
        for popup in driver.find_elements(By.XPATH, "//button[text()='X']"):
            if popup.is_displayed():
                driver.execute_script("arguments[0].click();", popup)
                time.sleep(1)
                break

        botao = wait.until(EC.presence_of_element_located((By.ID, "clinica-selecionada")))
        ActionChains(driver).move_to_element(botao).click(botao).perform()
        time.sleep(1)

        campo_busca = wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//input[@placeholder='Pesquisar clínica']")
        ))
        campo_busca.click()
        campo_busca.send_keys(CLINICA)
        time.sleep(2)

        opcao = wait.until(EC.presence_of_element_located(
            (By.XPATH, f"//a[contains(@class,'trocar-clinica') and contains(.,'{CLINICA}')]")
        ))
        driver.execute_script("arguments[0].click();", opcao)
        time.sleep(5)

        driver.get(f"{URL_BASE}/pacientes")
        time.sleep(3)
    except Exception:
        driver.quit()
        raise

    return driver, wait


def processar_paciente(driver, wait, pg, dados):
    """Localiza o paciente e cadastra os planos ainda não feitos (marca 'feito' em cada um)."""
    nome_paciente = dados["nomePaciente"]

    # 1. Localiza o paciente pelo PG
    try:
        url_paciente = localizar_paciente(driver, wait, pg)
        print(f"  Paciente localizado: {url_paciente}")
    except Exception as e:
        print(f"  ERRO ao localizar '{nome_paciente}' (PG: {pg}): {e}")
        salvar_print_erro(driver, f"{pg}_localizar")
        try:
            marcar_status(dados["linhas"], "erro")
            print("  Marcado como erro na planilha.")
        except Exception as erro_gravacao:
            print(f"  Falha ao gravar erro na planilha: {erro_gravacao}")
        return

    # 2. Cadastra cada plano (um por Perfil Tratamento)
    for perfil, plano in dados["planos"].items():
        if plano.get("feito"):
            continue
        print(f"  Cadastrando plano '{perfil}'...")

        try:
            cadastrar_plano(driver, wait, url_paciente, perfil, plano)
            marcar_status(plano["linhas"], "cadastrado")
            print(f"  Plano '{perfil}' cadastrado.")

        except Exception as e:
            print(f"  ERRO no plano '{perfil}' de '{nome_paciente}' (PG: {pg}): {e}")
            salvar_print_erro(driver, f"{pg}_{perfil}")
            try:
                marcar_status(plano["linhas"], "erro")
                print("  Marcado como erro na planilha.")
            except Exception as erro_gravacao:
                print(f"  Falha ao gravar erro na planilha: {erro_gravacao}")

        plano["feito"] = True


def navegador(fila, progresso, parar, atraso):
    """Um navegador: pega pacientes da fila até ela esvaziar.

    Se a sessão cair (login, Chrome fechado...), abre outra em 10s e devolve
    à fila o paciente que estava no meio, só com os planos que faltavam.
    """
    if parar.wait(atraso):  # escalona os logins para não entrarem todos juntos
        return

    while not parar.is_set() and not fila.empty():
        driver = None
        atual = None
        try:
            driver, wait = abrir_sessao()
            print("Sessão aberta.")

            while not parar.is_set():
                try:
                    atual = fila.get_nowait()
                except queue.Empty:
                    return

                pg, dados = atual
                with trava_progresso:
                    progresso["iniciados"] += 1
                    numero = progresso["iniciados"]
                print(f"\n[{numero}/{progresso['total']}] {'─' * 50}")
                imprimir_paciente(pg, dados)

                processar_paciente(driver, wait, pg, dados)
                atual = None
                with trava_progresso:
                    progresso["concluidos"] += 1

        except Exception as e:
            print(f"\nERRO na sessão: {resumo_erro(e)}")
            if driver:
                nome = threading.current_thread().name
                salvar_print_erro(driver, f"sessao_{nome}")
            if atual:
                fila.put(atual)
                with trava_progresso:
                    progresso["iniciados"] -= 1
            print("Reabrindo o navegador em 10 segundos...")
            parar.wait(10)

        finally:
            if driver:
                try:
                    driver.quit()
                except Exception:
                    pass


def executar():
    pacientes = carregar_pacientes()

    if LIMITE_PACIENTES:
        pacientes = dict(list(pacientes.items())[:LIMITE_PACIENTES])

    if not pacientes:
        print("Nenhum paciente pendente encontrado.")
        return

    total_planos = sum(len(d["planos"]) for d in pacientes.values())
    print(f"{len(pacientes)} paciente(s) pendente(s), {total_planos} plano(s) de tratamento.")

    if MODO_REVISAO:
        imprimir_amostra(pacientes, quantidade=len(pacientes))
        return

    # Cada paciente entra uma vez na fila: nunca dois navegadores no mesmo paciente
    fila = queue.Queue()
    for item in pacientes.items():
        fila.put(item)

    total = len(pacientes)
    quantidade = min(NUM_NAVEGADORES, total)
    print(f"Abrindo {quantidade} navegador(es).")

    progresso = {"concluidos": 0, "iniciados": 0, "total": total, "inicio": time.time()}
    parar = threading.Event()
    parar_relogio = threading.Event()
    threading.Thread(target=relogio_previsao, args=(progresso, parar_relogio), daemon=True).start()

    threads = [
        threading.Thread(
            target=navegador,
            args=(fila, progresso, parar, i * INTERVALO_ENTRE_LOGINS),
            name=f"N{i + 1}",
            daemon=True,
        )
        for i in range(quantidade)
    ]
    for t in threads:
        t.start()

    try:
        while any(t.is_alive() for t in threads):
            time.sleep(1)
    except KeyboardInterrupt:
        print("\nInterrompido: cada navegador termina o paciente atual e fecha "
              "(Ctrl+C de novo para sair na hora).")
        parar.set()
        for t in threads:
            t.join()
    finally:
        parar_relogio.set()

    if fila.empty() and not parar.is_set():
        print(f"\nCadastro de planos concluído! Tempo total: {formatar_duracao(time.time() - progresso['inicio'])}")


if __name__ == "__main__":
    executar()
