"""
Cadastro de orçamentos no Clínica nas Nuvens a partir de orcamentos.xlsx (aba Orcamentos).

Estrutura:
    paciente (CPF) -> 1 orçamento -> itens
    item = procedimento + valor unitário + quantidade

Cada linha da planilha é uma execução de procedimento. As linhas do mesmo paciente
com o mesmo procedimento E o mesmo valor viram um item; a quantidade é o número
de linhas. O mesmo procedimento com valores diferentes vira itens separados.

Status: gravado na coluna "Status" da própria planilha, em todas as linhas do
paciente. Pacientes com Status são pulados na próxima execução.
    cadastrado -> orçamento salvo e confirmado
    erro       -> falhou antes de salvar (nada foi criado); apague o Status para tentar de novo
    verificar  -> o Salvar final foi clicado, mas a página não confirmou: confira no
                  sistema se o orçamento existe antes de apagar o Status (evita duplicar)
Pacientes sem CPF são marcados 'erro' e não são cadastrados.

Passos na tela (por paciente):
    /orcamento/lista > Adicionar orçamento > filtro CPF/CNPJ > CPF no campo paciente
    > descrição + observações gerais (resumo dos itens) > para cada item: Adicionar >
    Adicionar procedimento > procedimento, "procedimento migrado", quantidade, valor >
    Salvar > Adicionar parcelas > situação Aprovado > desmarca "gera financeiro" > Salvar
"""

import builtins
import csv
import openpyxl
import os
import queue
import re
import time
import threading
import unicodedata
from collections import OrderedDict, defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import Select, WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.common.keys import Keys


PASTA = Path(__file__).resolve().parent
PASTA_ERROS = PASTA / "erros"
ARQUIVO = PASTA / "orcamentos.xlsx"
ABA = "Orcamentos"
CLINICA = "SOU COLUNA RIO VERDE - Clínica da Coluna AMS"
URL_BASE = "https://app.clinicanasnuvens.com.br"

# Login no Clínica nas Nuvens
USUARIO = "mmerlo++++migracao@soucolunarioverde.com.br"
SENHA = "CNN@foguete"

LIMITE_OBSERVACOES = 1000          # observações gerais do orçamento
LIMITE_DESCRICAO = 260             # descrição do orçamento (campo obrigatório)
DESCRICAO_ITEM = "procedimento migrado"

# Opções do autocomplete de paciente que NÃO são pacientes (o campo permite criar um novo)
OPCOES_IGNORADAS = ("cadastrar", "criar", "novo paciente", "adicionar", "carregar mais", "mais registros")

COL_CPF = "CPF do Paciente"
COL_PACIENTE = "Paciente"
COL_PRONTUARIO = "Prontuário"
COL_PROCEDIMENTO = "Procedimento"
COL_VALOR = "Valor"
COL_STATUS = "Status"
STATUS_FINALIZADOS = ("cadastrado", "erro", "verificar")

# True = só mostra os orçamentos montados, sem abrir o navegador nem gravar Status
MODO_REVISAO = False

# Quantos pacientes processar nesta execução (None = todos). Use 1 para testar.
LIMITE_PACIENTES = None

# Quantos navegadores cadastram ao mesmo tempo (cada Chrome usa ~0,5–1 GB de RAM)
NUM_NAVEGADORES = 1
# Segundos entre o login de um navegador e o do próximo
INTERVALO_ENTRE_LOGINS = 15

HEADLESS = True

# Logs da execução (pasta logs/ ao lado do script):
#   orcamentos_<data_hora>.log  -> tudo o que aparece no terminal
#   resultados_<data_hora>.csv  -> uma linha por paciente (abre no Excel)
PASTA_LOGS = PASTA / "logs"
CARIMBO_EXECUCAO = f"{datetime.now():%Y%m%d_%H%M%S}"
ARQUIVO_LOG = PASTA_LOGS / f"orcamentos_{CARIMBO_EXECUCAO}.log"
ARQUIVO_RESULTADOS = PASTA_LOGS / f"resultados_{CARIMBO_EXECUCAO}.csv"
LOG_ATIVO = False  # ligado em executar() (o modo revisão não gera arquivos)

# Só um navegador por vez mexe no progresso; e as linhas do terminal não se misturam
trava_print = threading.Lock()
trava_progresso = threading.Lock()


def print(*args, sep=" ", end="\n", **kwargs):
    """print com hora e nome do navegador ([N1], [N2]...) em cada linha; também grava no log."""
    nome = threading.current_thread().name
    prefixo = f"{datetime.now():%H:%M:%S} " + (f"[{nome}] " if re.fullmatch(r"N\d+", nome) else "")
    mensagem = sep.join(str(a) for a in args)
    mensagem = "\n".join(f"{prefixo}{linha}" if linha.strip() else linha for linha in mensagem.split("\n"))
    with trava_print:
        builtins.print(mensagem, end=end, **kwargs)
        if LOG_ATIVO:
            try:
                PASTA_LOGS.mkdir(exist_ok=True)
                with open(ARQUIVO_LOG, "a", encoding="utf-8") as f:
                    f.write(mensagem + end)
            except Exception:
                pass  # falha no arquivo de log não pode parar o cadastro


# ── Dados ──────────────────────────────────────────────────────────────────────

def texto(valor):
    return str(valor).strip() if valor is not None else ""


def normalizar(valor):
    """Minúsculas, sem acentos e com espaços simples — para comparar nomes."""
    valor = unicodedata.normalize("NFKD", texto(valor))
    valor = "".join(c for c in valor if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", valor).strip().lower()


def ler_valor(valor):
    """'194.44' / '194,44' / '' -> float."""
    valor = texto(valor).replace("R$", "").strip()
    if not valor:
        return 0.0
    if "," in valor:
        valor = valor.replace(".", "").replace(",", ".")
    return float(valor)


def formatar_valor(valor):
    return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def texto_celula(valor):
    """Célula -> texto: números inteiros sem '.0' (ex.: CPF digitado como número)."""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return texto(valor)


def ler_planilha():
    """Linhas da aba ABA como dicts (chave = nome da coluna) com '_linha' = nº da linha."""
    wb = openpyxl.load_workbook(ARQUIVO, read_only=True, data_only=True)
    linhas = wb[ABA].iter_rows(values_only=True)
    cabecalho = [texto(c) for c in next(linhas)]
    resultado = [dict(zip(cabecalho, valores), _linha=n) for n, valores in enumerate(linhas, start=2)]
    wb.close()
    return resultado


def carregar_orcamentos():
    """Monta os orçamentos pendentes:

    pacientes[CPF] = {
        "nome": str, "prontuario": str,
        "linhas": [nº das linhas na planilha],
        "itens": [{"procedimento": str, "valor": float, "quantidade": int, "linhas": [...]}],
    }

    Também devolve os pacientes sem CPF: [{"nome", "prontuario", "linhas"}].
    """
    pacientes = OrderedDict()
    sem_cpf = OrderedDict()

    for linha in ler_planilha():
        num_linha = linha["_linha"]
        procedimento = texto(linha.get(COL_PROCEDIMENTO))
        nome = texto(linha.get(COL_PACIENTE))

        # Linha vazia ou cabeçalho repetido no meio do arquivo
        if not procedimento or procedimento == COL_PROCEDIMENTO:
            continue
        if texto(linha.get(COL_STATUS)).lower() in STATUS_FINALIZADOS:
            continue

        cpf = re.sub(r"\D", "", texto_celula(linha.get(COL_CPF)))
        prontuario = texto_celula(linha.get(COL_PRONTUARIO))

        if not cpf:
            chave = prontuario or nome
            registro = sem_cpf.setdefault(chave, {"nome": nome, "prontuario": prontuario, "linhas": []})
            registro["linhas"].append(num_linha)
            continue
        cpf = cpf.zfill(11)

        paciente = pacientes.setdefault(cpf, {"nome": nome, "prontuario": prontuario, "linhas": [], "itens": OrderedDict()})
        paciente["linhas"].append(num_linha)

        valor = ler_valor(linha.get(COL_VALOR))
        item = paciente["itens"].setdefault(
            (procedimento, valor),
            {"procedimento": procedimento, "valor": valor, "quantidade": 0, "linhas": []},
        )
        item["quantidade"] += 1
        item["linhas"].append(num_linha)

    for paciente in pacientes.values():
        paciente["itens"] = list(paciente["itens"].values())

    return pacientes, list(sem_cpf.values())


_trava_planilha = threading.Lock()


def marcar_status(numeros_linha, valor):
    """Grava `valor` na coluna Status (criada se não existir) das linhas informadas.

    Salva numa cópia e só depois troca pelo original: se o script for interrompido
    no meio do salvamento, a planilha original fica intacta.
    """
    with _trava_planilha:
        wb = openpyxl.load_workbook(ARQUIVO)
        ws = wb[ABA]

        cabecalho = [texto(c.value) for c in ws[1]]
        if COL_STATUS in cabecalho:
            col_status = cabecalho.index(COL_STATUS) + 1
        else:
            col_status = len(cabecalho) + 1
            ws.cell(row=1, column=col_status, value=COL_STATUS)

        for num_linha in numeros_linha:
            ws.cell(row=num_linha, column=col_status, value=valor)

        temporario = ARQUIVO.with_name(f"~tmp_{ARQUIVO.name}")
        wb.save(temporario)
        os.replace(temporario, ARQUIVO)


def total_orcamento(dados):
    return sum(i["valor"] * i["quantidade"] for i in dados["itens"])


def imprimir_paciente(cpf, dados):
    print(f"  Paciente: {dados['nome']} (CPF: {cpf}, prontuário {dados['prontuario']}) | "
          f"{len(dados['linhas'])} linha(s) na planilha")
    for item in dados["itens"]:
        print(f"      {item['quantidade']:3d}x {item['procedimento']} — {formatar_valor(item['valor'])} cada")
    print(f"      Total do orçamento: {formatar_valor(total_orcamento(dados))}")


def formatar_duracao(segundos):
    segundos = int(segundos)
    horas, resto = divmod(segundos, 3600)
    minutos, segundos = divmod(resto, 60)
    return f"{horas:02d}:{minutos:02d}:{segundos:02d}"


def placar():
    return (f"✔ {CONTAGEM['cadastrado']} cadastrado(s) · ✖ {CONTAGEM['erro']} erro(s) · "
            f"⚠ {CONTAGEM['verificar']} verificar")


def relogio_previsao(progresso, parar, intervalo=60):
    """A cada `intervalo` segundos, imprime o progresso, o placar e a previsão de término."""
    while not parar.wait(intervalo):
        concluidos = progresso["concluidos"]
        total = progresso["total"]
        decorrido = time.time() - progresso["inicio"]

        if concluidos == 0:
            print(f"\n⏱  [{formatar_duracao(decorrido)} decorrido] 0/{total} | {placar()} | calculando previsão...")
            continue

        media = decorrido / concluidos
        restante = media * (total - concluidos)
        termino = datetime.now() + timedelta(seconds=restante)
        print(
            f"\n⏱  {concluidos}/{total} pacientes ({concluidos / total:.0%}) | {placar()}\n"
            f"⏱  média {media:.0f}s/paciente (somando os navegadores) | "
            f"restante: {formatar_duracao(restante)} | término previsto: {termino:%d/%m %H:%M}"
        )


# ── Relatório de resultados ────────────────────────────────────────────────────

RESULTADOS = []
CONTAGEM = defaultdict(int)
trava_resultados = threading.Lock()
COLUNAS_RESULTADOS = ["data_hora", "navegador", "cpf", "paciente", "prontuario", "itens",
                      "total", "status", "segundos", "detalhe"]


def registrar_resultado(cpf, dados, status, detalhe="", inicio=None):
    """Guarda o resultado do paciente (para o resumo final) e acrescenta no CSV de resultados."""
    registro = {
        "data_hora": f"{datetime.now():%d/%m/%Y %H:%M:%S}",
        "navegador": threading.current_thread().name,
        "cpf": cpf,
        "paciente": dados.get("nome", ""),
        "prontuario": dados.get("prontuario", ""),
        "itens": len(dados.get("itens", [])),
        "total": f"{total_orcamento(dados):.2f}".replace(".", ",") if dados.get("itens") else "",
        "status": status,
        "segundos": f"{time.time() - inicio:.0f}" if inicio else "",
        "detalhe": detalhe,
    }
    with trava_resultados:
        RESULTADOS.append(registro)
        CONTAGEM[status] += 1
        if not LOG_ATIVO:
            return
        try:
            PASTA_LOGS.mkdir(exist_ok=True)
            novo = not ARQUIVO_RESULTADOS.exists()
            # ';' e utf-8 com BOM: abre direto no Excel em português
            with open(ARQUIVO_RESULTADOS, "a", newline="", encoding="utf-8-sig") as f:
                escritor = csv.DictWriter(f, fieldnames=COLUNAS_RESULTADOS, delimiter=";")
                if novo:
                    escritor.writeheader()
                escritor.writerow(registro)
        except Exception as e:
            builtins.print(f"(não consegui gravar o CSV de resultados: {e})")


def imprimir_resumo(progresso, fila):
    duracao = formatar_duracao(time.time() - progresso["inicio"])
    print("\n" + "═" * 70)
    print(f"RESUMO — {len(RESULTADOS)} paciente(s) processado(s) em {duracao}")
    print(placar())
    if fila is not None and not fila.empty():
        print(f"Ainda na fila (não processados): {fila.qsize()} — rode de novo para continuar.")

    for status, titulo in (("erro", "ERROS (nada foi salvo; apague o Status para tentar de novo)"),
                           ("verificar", "VERIFICAR NO SISTEMA (o Salvar foi clicado mas não confirmou)")):
        lista = [r for r in RESULTADOS if r["status"] == status]
        if lista:
            print(f"\n{titulo}:")
            for r in lista:
                print(f"  - {r['paciente'] or '(sem nome)'} (CPF {r['cpf'] or '-'}, prontuário {r['prontuario'] or '-'}): {r['detalhe']}")

    if LOG_ATIVO:
        print(f"\nLog completo:   {ARQUIVO_LOG}")
        print(f"Resultados CSV: {ARQUIVO_RESULTADOS}")
    print("═" * 70)


def salvar_print_erro(driver, nome):
    try:
        PASTA_ERROS.mkdir(exist_ok=True)
        nome = re.sub(r"[^\w\-]+", "_", nome)
        driver.save_screenshot(str(PASTA_ERROS / f"{nome}.png"))
        print(f"  Screenshot salvo em {PASTA_ERROS / (nome + '.png')}")
    except Exception:
        pass


# ── Passos na plataforma ───────────────────────────────────────────────────────

def clicar(driver, elemento):
    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", elemento)
    driver.execute_script("arguments[0].click();", elemento)


def clicar_mouse(driver, elemento):
    """Clique 'real' (mousedown/mouseup/click) — necessário para select2 e autocomplete."""
    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", elemento)
    ActionChains(driver).move_to_element(elemento).pause(0.2).click(elemento).perform()


def resumo_erro(e):
    """1ª linha da mensagem (as do Selenium trazem um stacktrace enorme), com o tipo do
    erro só quando ele diz algo (ex.: TimeoutException) — não para Exception genérica."""
    mensagem = texto(getattr(e, "msg", None) or e).split("\n")[0]
    if type(e) is Exception:
        return mensagem or "erro sem mensagem"
    return f"{type(e).__name__}: {mensagem}" if mensagem else type(e).__name__


def digitar_no_campo(driver, campo, valor):
    """Foca, limpa e digita no campo; se o clique for bloqueado (algo por cima), foca via JS."""
    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", campo)
    try:
        campo.click()
    except Exception:
        driver.execute_script("arguments[0].focus();", campo)
    campo.clear()
    campo.send_keys(valor)


def itens_autocomplete(driver):
    return [
        li for li in driver.find_elements(By.CSS_SELECTOR, "ul.ui-autocomplete li.ui-menu-item")
        if li.is_displayed()
    ]


def selecionar_autocomplete(driver, campo, nome, preferir_exato=True, tentativas=3, exigir_nome=False, ignorar=()):
    """Digita `nome` num campo autocomplete e escolhe a opção.

    Com `preferir_exato`, prefere a opção com o nome exato; senão pega a primeira.
    Com `exigir_nome`, só aceita opção que contenha o nome digitado (senão falha).
    `ignorar`: trechos de texto de opções que não são resultados (ex.: "Cadastrar novo").
    Tenta primeiro com clique do mouse e, se não 'grudar', pelo teclado
    (seta para baixo + Enter).
    """
    for tentativa in range(1, tentativas + 1):
        try:
            digitar_no_campo(driver, campo, nome)
        except Exception as e:
            print(f"    Tentativa {tentativa}: não consegui digitar '{nome}' no campo ({resumo_erro(e)}).")
            time.sleep(2)
            continue

        try:
            WebDriverWait(driver, 20).until(lambda d: itens_autocomplete(d) or False)
        except TimeoutException:
            print(f"    Tentativa {tentativa}: nenhuma opção apareceu para '{nome}'.")
            continue
        time.sleep(1)  # espera a lista terminar de carregar

        itens = itens_autocomplete(driver)
        if not itens:
            continue

        textos_todos = [texto(li.text).split("\n")[0] for li in itens]
        textos = [t for t in textos_todos if not any(p in normalizar(t) for p in ignorar)]
        if not textos:
            raise Exception(f"nenhum resultado para '{nome}' (opções: {textos_todos[:5]})")
        # Posição de cada opção válida na lista completa (para clique e teclado)
        indices = [i for i, t in enumerate(textos_todos) if not any(p in normalizar(t) for p in ignorar)]
        itens = [itens[i] for i in indices]
        posicao = 0
        if exigir_nome:
            compativeis = [i for i, t in enumerate(textos) if normalizar(nome) in normalizar(t)]
            if not compativeis:
                raise Exception(f"nenhuma opção corresponde a '{nome}' (opções: {textos[:5]})")
            exatos = [i for i in compativeis if normalizar(textos[i]) == normalizar(nome)]
            posicao = (exatos or compativeis)[0]
        elif preferir_exato:
            posicao = next((i for i, t in enumerate(textos) if normalizar(t) == normalizar(nome)), 0)
        escolhido = textos[posicao]

        try:
            if tentativa == 1:
                clicar_mouse(driver, itens[posicao])
            else:
                for _ in range(indices[posicao] + 1):
                    campo.send_keys(Keys.ARROW_DOWN)
                    time.sleep(0.2)
                campo.send_keys(Keys.RETURN)
        except Exception as e:
            print(f"    Tentativa {tentativa}: falha ao clicar na opção ({e}).")
            continue
        time.sleep(1)

        valor = texto(campo.get_attribute("value"))
        if valor and not itens_autocomplete(driver):
            return valor

        print(f"    Tentativa {tentativa}: opção '{escolhido}' não ficou selecionada.")

    raise Exception(f"não foi possível selecionar '{nome}'")


def abrir_sessao():
    """Abre o Chrome, faz login e seleciona a clínica. Devolve (driver, wait)."""
    opcoes = webdriver.ChromeOptions()
    if HEADLESS:
        opcoes.add_argument("--headless=new")
        opcoes.add_argument("--window-size=1920,1080")
    driver = webdriver.Chrome(options=opcoes)
    if not HEADLESS:
        driver.maximize_window()
    wait = WebDriverWait(driver, 15)

    try:
        # ── Login ──────────────────────────────────────────────────────────────
        driver.get(f"{URL_BASE}/")

        campo_email = wait.until(EC.presence_of_element_located(
            (By.XPATH, "/html/body/div[1]/div/div[2]/div[1]/div/div[3]/form/div[1]/input")
        ))
        campo_email.send_keys(USUARIO)
        driver.find_element(By.XPATH, "/html/body/div[1]/div/div[2]/div[1]/div/div[3]/form/div[2]/button").click()

        campo_senha = wait.until(EC.element_to_be_clickable((By.ID, "password")))
        campo_senha.click()
        campo_senha.send_keys(SENHA)
        driver.find_element(By.XPATH, "/html/body/div[1]/div/div[2]/div[1]/div/div[4]/form/div[4]/button").click()
        time.sleep(5)

        # ── Seleciona a clínica ────────────────────────────────────────────────
        for popup in driver.find_elements(By.XPATH, "//button[text()='X']"):
            if popup.is_displayed():
                driver.execute_script("arguments[0].click();", popup)
                time.sleep(1)
                break

        botao = wait.until(EC.presence_of_element_located((By.ID, "clinica-selecionada")))
        if clinica_atual(botao) == normalizar(CLINICA):
            print(f"Clínica '{CLINICA}' já selecionada.")
        else:
            print(f"Trocando a clínica '{botao.get_attribute('alt')}' por '{CLINICA}'...")
            ActionChains(driver).move_to_element(botao).click(botao).perform()
            time.sleep(1)

            campo_busca = wait.until(EC.element_to_be_clickable(
                (By.XPATH, "//input[@placeholder='Pesquisar clínica']")
            ))
            campo_busca.click()
            campo_busca.send_keys(CLINICA)
            time.sleep(2)

            opcao = wait.until(EC.presence_of_element_located(
                (By.XPATH, f"//a[contains(@class,'trocar-clinica') and contains(.,'{CLINICA}')]")
            ))
            driver.execute_script("arguments[0].click();", opcao)
            time.sleep(5)

            botao = wait.until(EC.presence_of_element_located((By.ID, "clinica-selecionada")))
            if clinica_atual(botao) != normalizar(CLINICA):
                raise Exception(f"não foi possível selecionar a clínica (ficou '{botao.get_attribute('alt')}')")
    except Exception:
        driver.quit()
        raise

    return driver, wait


def clinica_atual(botao):
    """Nome da clínica selecionada, lido do elemento #clinica-selecionada (atributo alt)."""
    return normalizar(botao.get_attribute("alt") or botao.get_attribute("data-bs-original-title") or botao.text)


# ── Cadastro do orçamento ──────────────────────────────────────────────────────

def tamanho_textarea(valor):
    # Por segurança, conta a quebra de linha como 2 caracteres (\r\n) no maxlength
    return len(valor) + valor.count("\n")


def cortar(textarea, limite):
    """Corta o texto com '…' para caber no maxlength do campo."""
    if tamanho_textarea(textarea) <= limite:
        return textarea
    while tamanho_textarea(textarea + "…") > limite:
        textarea = textarea[:-1]
    return textarea.rstrip() + "…"


def titulo_orcamento(cpf, dados):
    return f"Orçamento migrado — {dados['nome']} (CPF {cpf}, prontuário {dados['prontuario']})"


def montar_descricao(cpf, dados):
    """Descrição do orçamento: a 1ª linha das observações, cabendo em LIMITE_DESCRICAO."""
    return cortar(titulo_orcamento(cpf, dados), LIMITE_DESCRICAO)


def montar_observacoes(cpf, dados):
    """Resumo do orçamento para as observações gerais, cabendo em LIMITE_OBSERVACOES."""
    linhas = [titulo_orcamento(cpf, dados)]
    for item in dados["itens"]:
        linhas.append(f"{item['quantidade']}x {item['procedimento']} — {formatar_valor(item['valor'])} cada")
    linhas.append(f"Total: {formatar_valor(total_orcamento(dados))}")
    return cortar("\n".join(linhas), LIMITE_OBSERVACOES)


def visivel(driver, seletor, raiz=None, timeout=15):
    """Primeiro elemento VISÍVEL que casa com o seletor CSS (resolve IDs repetidos na página)."""
    raiz = raiz or driver

    def procurar(_):
        for e in raiz.find_elements(By.CSS_SELECTOR, seletor):
            try:
                if e.is_displayed():
                    return e
            except Exception:
                continue
        return False

    try:
        return WebDriverWait(driver, timeout).until(procurar)
    except TimeoutException:
        raise Exception(f"elemento '{seletor}' não apareceu na tela")


def selecionar_opcao(driver, elemento_select, valor):
    """<select> pelo value; confere e, se não pegou, seleciona via jQuery/JS e dispara o change."""
    try:
        Select(elemento_select).select_by_value(valor)
    except Exception:
        pass
    if elemento_select.get_attribute("value") != valor:
        driver.execute_script("""
            var s = arguments[0];
            if (window.jQuery) { jQuery(s).val(arguments[1]).trigger('change'); }
            else { s.value = arguments[1]; s.dispatchEvent(new Event('change', {bubbles: true})); }
        """, elemento_select, valor)
    else:
        driver.execute_script("arguments[0].dispatchEvent(new Event('change', {bubbles: true}));", elemento_select)
    time.sleep(0.5)
    if elemento_select.get_attribute("value") != valor:
        raise Exception(f"não foi possível selecionar '{valor}' (ficou '{elemento_select.get_attribute('value')}')")


def ler_numero(valor):
    """'1.234,56' / '1,234.56' / 'R$ 100,00' -> float (o último separador é o decimal)."""
    valor = re.sub(r"[^\d,.]", "", texto(valor))
    if not valor:
        return None
    posicao = max(valor.rfind(","), valor.rfind("."))
    if posicao == -1:
        return float(valor)
    inteiro = re.sub(r"\D", "", valor[:posicao]) or "0"
    decimal = re.sub(r"\D", "", valor[posicao + 1:])
    return float(f"{inteiro}.{decimal}")


def preencher_numero(driver, campo, numero):
    """Campo numérico ou com máscara de dinheiro: tenta formas de digitar até o valor
    exibido bater com `numero` (máscaras às vezes tratam os dígitos como centavos)."""
    alvo = round(float(numero), 2)
    inteiro = int(alvo) if alvo.is_integer() else alvo
    tentativas = [f"{alvo:.2f}".replace(".", ","), f"{alvo:.2f}".replace(".", ""), f"{inteiro}", f"{alvo:.2f}"]
    for digitado in dict.fromkeys(tentativas):
        try:
            campo.click()
        except Exception:
            driver.execute_script("arguments[0].focus();", campo)
        driver.execute_script("arguments[0].select();", campo)
        campo.send_keys(Keys.BACKSPACE)
        campo.send_keys(digitado)
        time.sleep(0.3)
        if ler_numero(campo.get_attribute("value")) == alvo:
            return

    # Plano B: grava direto e dispara os eventos
    driver.execute_script("""
        arguments[0].value = arguments[1];
        arguments[0].dispatchEvent(new Event('input', {bubbles: true}));
        arguments[0].dispatchEvent(new Event('change', {bubbles: true}));
    """, campo, f"{alvo:.2f}".replace(".", ","))
    time.sleep(0.3)
    if ler_numero(campo.get_attribute("value")) != alvo:
        raise Exception(f"não foi possível preencher {numero} (campo ficou '{campo.get_attribute('value')}')")


def preencher_inteiro(driver, campo, numero):
    """Campo de quantidade (type=number): seleciona tudo, digita o inteiro e confere."""
    valor = str(int(numero))
    try:
        campo.click()
    except Exception:
        driver.execute_script("arguments[0].focus();", campo)
    driver.execute_script("arguments[0].select();", campo)
    campo.send_keys(Keys.BACKSPACE)
    campo.send_keys(valor)
    time.sleep(0.3)

    # Plano B: grava direto no campo e dispara os eventos
    if texto(campo.get_attribute("value")) != valor:
        driver.execute_script("""
            arguments[0].value = arguments[1];
            arguments[0].dispatchEvent(new Event('input', {bubbles: true}));
            arguments[0].dispatchEvent(new Event('change', {bubbles: true}));
        """, campo, valor)
        time.sleep(0.3)

    if texto(campo.get_attribute("value")) != valor:
        raise Exception(f"não foi possível preencher a quantidade {valor} (campo ficou '{campo.get_attribute('value')}')")


def abrir_novo_orcamento(driver, wait):
    """/orcamento/lista > Adicionar orçamento."""
    driver.get(f"{URL_BASE}/orcamento/lista")
    botao_novo = wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, "a[href='/orcamento/novo']")))
    clicar(driver, botao_novo)
    visivel(driver, "#paciente")
    time.sleep(1)


def mesmo_paciente(nome_planilha, nome_sistema):
    """Confere o paciente escolhido: nomes iguais (sem acento/maiúscula), um contido no
    outro, ou mesmo primeiro e último nome."""
    a, b = normalizar(nome_planilha), normalizar(nome_sistema)
    if not a or not b:
        return False
    if a == b or a in b or b in a:
        return True
    pa, pb = a.split(), b.split()
    return pa[0] == pb[0] and pa[-1] in pb


def selecionar_paciente(driver, cpf, nome):
    """Filtro CPF/CNPJ > digita o CPF > escolhe o primeiro paciente > confere o nome."""
    selecionar_opcao(driver, visivel(driver, "#filtro-paciente"), "CPF_CNPJ")

    campo = visivel(driver, "#paciente")
    escolhido = selecionar_autocomplete(driver, campo, cpf, preferir_exato=False, ignorar=OPCOES_IGNORADAS)
    if not mesmo_paciente(nome, escolhido):
        raise Exception(f"paciente selecionado '{escolhido}' não confere com '{nome}' da planilha")
    print(f"    Paciente selecionado: {escolhido}")


def preencher_texto(driver, seletor, valor, nome_campo):
    """Preenche um campo de texto (textarea) e confere que ficou preenchido."""
    campo = visivel(driver, seletor)
    campo.clear()
    campo.send_keys(valor)
    if not texto(campo.get_attribute("value")):
        raise Exception(f"{nome_campo} não foi preenchida")


def adicionar_item(driver, wait, item):
    """Adicionar > Adicionar procedimento > procedimento, descrição, quantidade, valor > Salvar."""
    botao_adicionar = wait.until(EC.element_to_be_clickable(
        (By.XPATH, "//button[@data-kt-menu-trigger='click' and contains(normalize-space(.),'Adicionar')]")
    ))
    clicar_mouse(driver, botao_adicionar)

    # Espera o menu abrir; se o clique do mouse não abriu, tenta via JS
    xpath_opcao = "//span[contains(@class,'menu-title') and normalize-space()='Adicionar procedimento']"
    try:
        opcao = WebDriverWait(driver, 3).until(EC.visibility_of_element_located((By.XPATH, xpath_opcao)))
    except TimeoutException:
        clicar(driver, botao_adicionar)
        time.sleep(0.5)
        opcao = driver.find_element(By.XPATH, xpath_opcao)
    # O clique vai no item do menu (pai do texto), que é quem tem o evento
    alvo = opcao.find_elements(By.XPATH, "./ancestor::*[contains(@class,'menu-link')][1]")
    clicar(driver, alvo[0] if alvo else opcao)

    campo_proc = visivel(driver, "#campo_proc")
    caixa = container_do_campo(campo_proc)
    time.sleep(0.5)

    escolhido = selecionar_autocomplete(driver, campo_proc, item["procedimento"])
    if normalizar(escolhido) != normalizar(item["procedimento"]):
        print(f"    Atenção: procedimento '{escolhido}' selecionado para '{item['procedimento']}'.")
    time.sleep(1)  # o sistema preenche o valor da tabela depois da seleção

    campo_desc = visivel(driver, "#campo_desc", raiz=caixa)
    campo_desc.clear()
    campo_desc.send_keys(DESCRICAO_ITEM)

    campo_qtd = visivel(driver, "#campo_qtd", raiz=caixa)
    preencher_inteiro(driver, campo_qtd, item["quantidade"])

    campo_valor = visivel(driver, "#campo_valor", raiz=caixa)
    preencher_numero(driver, campo_valor, item["valor"])

    # Confere de novo logo antes de salvar (o JS da tela pode recalcular os campos)
    if ler_numero(campo_qtd.get_attribute("value")) != float(item["quantidade"]):
        raise Exception(f"quantidade mudou para '{campo_qtd.get_attribute('value')}' antes de salvar")
    if ler_numero(campo_valor.get_attribute("value")) != round(item["valor"], 2):
        raise Exception(f"valor mudou para '{campo_valor.get_attribute('value')}' antes de salvar")

    botoes = [
        b for b in caixa.find_elements(By.CSS_SELECTOR, "button[type='submit']")
        if b.is_displayed() and "Salvar" in b.text
    ]
    if not botoes:
        raise Exception("botão Salvar do procedimento não encontrado")
    clicar(driver, botoes[-1])

    # A caixa fecha quando o procedimento é salvo
    try:
        WebDriverWait(driver, 10).until(lambda d: not campo_proc.is_displayed())
    except TimeoutException:
        raise Exception(f"a caixa do procedimento '{item['procedimento']}' não fechou após salvar (erro de validação?)")
    except Exception:
        pass  # campo removido da página = salvo
    time.sleep(0.5)


def container_do_campo(campo):
    """Caixa (modal) onde está o campo do procedimento; se não for modal, o form."""
    for xpath in ("./ancestor::div[contains(concat(' ', @class, ' '), ' modal ')][1]", "./ancestor::form[1]"):
        encontrados = campo.find_elements(By.XPATH, xpath)
        if encontrados:
            return encontrados[0]
    raise Exception("caixa do procedimento não encontrada")


def gerar_parcelas(driver):
    """Clica em 'Adicionar parcelas' (parcelamento.gerarParcelas()) depois dos procedimentos."""
    botao = visivel(
        driver,
        "button[onclick*='gerarParcelas']",
        timeout=10,
    )
    clicar(driver, botao)
    confirmar_alerta_se_houver(driver)
    time.sleep(1.5)  # as parcelas são montadas por JS na própria tela
    print("    Parcelas geradas.")


def aprovar_sem_financeiro(driver):
    """Situação = Aprovado e desmarca 'gera financeiro'."""
    selecionar_opcao(driver, visivel(driver, "#orcamento-situacao"), "APROVADO")

    # O checkbox pode aparecer só depois de escolher Aprovado
    try:
        checkbox = WebDriverWait(driver, 5).until(
            lambda d: d.find_elements(By.CSS_SELECTOR, "input[type='checkbox'][name='extra.geraFinanceiro']") or False
        )[0]
    except TimeoutException:
        raise Exception("opção 'gera financeiro' não encontrada — não salvei para não gerar financeiro")

    if checkbox.is_selected():
        clicar(driver, checkbox)
        time.sleep(0.3)
    if checkbox.is_selected():
        driver.execute_script("""
            arguments[0].checked = false;
            arguments[0].dispatchEvent(new Event('change', {bubbles: true}));
        """, checkbox)
        time.sleep(0.3)
    if checkbox.is_selected():
        raise Exception("não foi possível desmarcar 'gera financeiro'")


def confirmar_alerta_se_houver(driver, timeout=2):
    """Se aparecer um alerta de confirmação (SweetAlert), clica em confirmar."""
    try:
        botao = visivel(driver, ".swal2-confirm", timeout=timeout)
        clicar(driver, botao)
        time.sleep(1)
    except Exception:
        pass


def salvar_orcamento(driver, dados):
    """Salvar final. Marca dados['salvar_clicado'] antes do clique: se a confirmação não
    vier, o paciente fica 'verificar' em vez de ser refeito (evita orçamento duplicado)."""
    botoes = [
        b for b in driver.find_elements(
            By.XPATH,
            "//button[@type='submit' and contains(normalize-space(.),'Salvar')"
            " and not(ancestor::div[contains(concat(' ', @class, ' '), ' modal ')])]",
        )
        if b.is_displayed()
    ]
    if not botoes:
        raise Exception("botão Salvar do orçamento não encontrado")

    url_antes = driver.current_url
    dados["salvar_clicado"] = True
    clicar(driver, botoes[-1])
    confirmar_alerta_se_houver(driver)

    try:
        WebDriverWait(driver, 30).until(lambda d: d.current_url != url_antes)
    except TimeoutException:
        raise Exception("o orçamento não confirmou o salvamento (a página não saiu do cadastro)")
    time.sleep(1)


def cadastrar_orcamento(driver, wait, cpf, dados):
    abrir_novo_orcamento(driver, wait)
    selecionar_paciente(driver, cpf, dados["nome"])
    preencher_texto(driver, "#descricaoOrcamento", montar_descricao(cpf, dados), "descrição do orçamento")
    preencher_texto(driver, "#orcamento-observacoesGerais", montar_observacoes(cpf, dados), "observações gerais")

    total_itens = len(dados["itens"])
    for n, item in enumerate(dados["itens"], start=1):
        print(f"    Item {n}/{total_itens}: {item['quantidade']}x {item['procedimento']} — {formatar_valor(item['valor'])}")
        try:
            adicionar_item(driver, wait, item)
        except Exception as e:
            raise Exception(f"item {n}/{total_itens} '{item['procedimento']}': {resumo_erro(e)}") from e

    gerar_parcelas(driver)
    aprovar_sem_financeiro(driver)
    print("    Situação Aprovado, gera financeiro desmarcado.")
    print("    Salvando orçamento...")
    salvar_orcamento(driver, dados)


def sessao_ativa(driver):
    """False se o Chrome fechou ou a sessão caiu (voltou para a tela de login)."""
    try:
        driver.current_url
        return not any(e.is_displayed() for e in driver.find_elements(By.ID, "password"))
    except Exception:
        return False


# Quedas de sessão por paciente: um paciente que derruba o navegador repetidamente
# vira 'erro' em vez de travar o lote.
MAX_QUEDAS_POR_PACIENTE = 2
QUEDAS = defaultdict(int)

# Sessões que não conseguem abrir em sequência (ex.: senha errada, site fora do ar)
MAX_FALHAS_DE_LOGIN = 5

# Sinal para todos os navegadores pararem (ex.: Status de um orçamento salvo não foi gravado)
PARAR = threading.Event()
MOTIVO_PARADA = []


class SessaoCaiu(Exception):
    """O Chrome fechou ou a sessão expirou no meio do paciente (não é erro do paciente)."""


class StatusNaoGravado(Exception):
    """O orçamento foi enviado, mas o Status não foi gravado: parar tudo para não duplicar."""


def gravar_status(numeros_linha, valor, tentativas=6):
    """marcar_status com novas tentativas (ex.: planilha aberta no Excel/LibreOffice).

    Se não conseguir gravar 'cadastrado'/'verificar', anota num arquivo à parte e para
    todos os navegadores: seguir em frente faria a próxima execução cadastrar o
    orçamento de novo.
    """
    for tentativa in range(1, tentativas + 1):
        try:
            marcar_status(numeros_linha, valor)
            return
        except Exception as e:
            print(f"  Não consegui gravar o Status na planilha ({resumo_erro(e)}) — tentativa {tentativa}/{tentativas}.")
            time.sleep(5)

    pendente = PASTA / "status_nao_gravado.txt"
    with open(pendente, "a", encoding="utf-8") as f:
        f.write(f"{datetime.now():%d/%m/%Y %H:%M:%S} | {valor} | linhas {numeros_linha}\n")
    if valor in ("cadastrado", "verificar"):
        mensagem = (
            f"PARADO: um orçamento foi enviado, mas o Status '{valor}' não foi gravado na planilha "
            f"(anotado em {pendente.name}). Feche a planilha no Excel/LibreOffice, grave o Status "
            "dessas linhas à mão e rode de novo."
        )
        MOTIVO_PARADA.append(mensagem)
        PARAR.set()
        raise StatusNaoGravado(mensagem)


# ── Execução ───────────────────────────────────────────────────────────────────

def processar_paciente(driver, wait, cpf, dados):
    """Cadastra o orçamento de um paciente e grava o Status.

    Levanta SessaoCaiu se o navegador morreu antes do Salvar final (o paciente volta
    para a fila); qualquer outro problema vira 'erro' ou 'verificar' na planilha.
    """
    dados["salvar_clicado"] = False
    inicio = time.time()
    try:
        cadastrar_orcamento(driver, wait, cpf, dados)
    except Exception as e:
        motivo = resumo_erro(e)
        print(f"  ✖ ERRO no orçamento de '{dados['nome']}' (CPF: {cpf}): {motivo}")
        salvar_print_erro(driver, f"{cpf}_orcamento")

        if dados.get("salvar_clicado"):
            # Pode ter salvado: não refaz, para não duplicar o orçamento
            gravar_status(dados["linhas"], "verificar")
            registrar_resultado(cpf, dados, "verificar", motivo, inicio)
            print("  ⚠ Marcado como VERIFICAR: confira no sistema se o orçamento foi criado.")
        elif not sessao_ativa(driver):
            raise SessaoCaiu(motivo)
        else:
            gravar_status(dados["linhas"], "erro")
            registrar_resultado(cpf, dados, "erro", motivo, inicio)
            print("  Marcado como erro na planilha (nada foi salvo para este paciente).")
        return

    gravar_status(dados["linhas"], "cadastrado")
    registrar_resultado(cpf, dados, "cadastrado", "", inicio)
    print(f"  ✔ Orçamento cadastrado em {time.time() - inicio:.0f}s ({total_orcamento_fmt(dados)}).")


def total_orcamento_fmt(dados):
    return f"{len(dados['itens'])} item(ns), total {formatar_valor(total_orcamento(dados))}"


def navegador(fila, progresso, atraso):
    """Um navegador: faz login e pega pacientes da fila até ela esvaziar.

    Se a sessão cair, reabre em 10s e devolve à fila o paciente que estava no meio
    (até MAX_QUEDAS_POR_PACIENTE vezes; depois ele vira 'erro').
    """
    if PARAR.wait(atraso):  # escalona os logins para não entrarem todos juntos
        return

    falhas_de_login = 0
    while not PARAR.is_set() and not fila.empty():
        driver = None
        atual = None
        try:
            driver, wait = abrir_sessao()
            falhas_de_login = 0
            print("Sessão aberta.")

            while not PARAR.is_set():
                try:
                    atual = fila.get_nowait()
                except queue.Empty:
                    return

                cpf, dados = atual
                with trava_progresso:
                    progresso["iniciados"] += 1
                    numero = progresso["iniciados"]
                print(f"\n[{numero}/{progresso['total']}] {'─' * 50}")
                imprimir_paciente(cpf, dados)

                processar_paciente(driver, wait, cpf, dados)
                atual = None
                with trava_progresso:
                    progresso["concluidos"] += 1

        except StatusNaoGravado:
            return

        except Exception as e:
            if atual is None:
                falhas_de_login += 1
                print(f"\nERRO ao abrir a sessão ({falhas_de_login}/{MAX_FALHAS_DE_LOGIN}): {resumo_erro(e)}")
                if falhas_de_login >= MAX_FALHAS_DE_LOGIN:
                    print("Desistindo deste navegador (login falhou várias vezes seguidas).")
                    return
            else:
                cpf, dados = atual
                QUEDAS[cpf] += 1
                print(f"\nSessão caiu no paciente {cpf} ({QUEDAS[cpf]}/{MAX_QUEDAS_POR_PACIENTE}): {resumo_erro(e)}")
                if driver:
                    salvar_print_erro(driver, f"sessao_{threading.current_thread().name}_{cpf}")
                with trava_progresso:
                    progresso["iniciados"] -= 1
                if QUEDAS[cpf] < MAX_QUEDAS_POR_PACIENTE:
                    fila.put(atual)
                    print("  Paciente devolvido à fila.")
                else:
                    try:
                        gravar_status(dados["linhas"], "erro")
                        registrar_resultado(cpf, dados, "erro",
                                            f"derrubou a sessão {QUEDAS[cpf]} vezes: {resumo_erro(e)}")
                        print("  Paciente derrubou a sessão várias vezes — marcado como erro.")
                    except StatusNaoGravado:
                        return
                    with trava_progresso:
                        progresso["concluidos"] += 1
            print("Reabrindo o navegador em 10 segundos...")
            PARAR.wait(10)

        finally:
            if driver:
                try:
                    driver.quit()
                except Exception:
                    pass


def executar():
    pacientes, sem_cpf = carregar_orcamentos()

    if LIMITE_PACIENTES:
        pacientes = OrderedDict(list(pacientes.items())[:LIMITE_PACIENTES])

    total_itens = sum(len(d["itens"]) for d in pacientes.values())
    print(f"{len(pacientes)} paciente(s) pendente(s), {total_itens} item(ns) de orçamento.")
    if sem_cpf:
        print(f"{len(sem_cpf)} paciente(s) sem CPF (serão marcados como erro):")
        for p in sem_cpf:
            print(f"  - {p['nome'] or '(sem nome)'} | prontuário {p['prontuario'] or '-'} | linhas {p['linhas']}")

    if MODO_REVISAO:
        for cpf, dados in pacientes.items():
            print()
            imprimir_paciente(cpf, dados)
        return

    global LOG_ATIVO
    LOG_ATIVO = True
    print(f"Início da execução | {len(pacientes)} paciente(s) | {NUM_NAVEGADORES} navegador(es) | "
          f"{'headless' if HEADLESS else 'com janela'} | clínica {CLINICA}")

    for p in sem_cpf:
        gravar_status(p["linhas"], "erro")
        registrar_resultado("", p, "erro", "paciente sem CPF na planilha")

    if not pacientes:
        print("Nenhum paciente pendente encontrado.")
        return

    # Cada paciente entra uma vez na fila: nunca dois navegadores no mesmo paciente
    fila = queue.Queue()
    for item in pacientes.items():
        fila.put(item)

    total = len(pacientes)
    quantidade = min(NUM_NAVEGADORES, total)
    print(f"Abrindo {quantidade} navegador(es).")

    progresso = {"concluidos": 0, "iniciados": 0, "total": total, "inicio": time.time()}
    parar_relogio = threading.Event()
    threading.Thread(target=relogio_previsao, args=(progresso, parar_relogio), daemon=True).start()

    threads = [
        threading.Thread(
            target=navegador,
            args=(fila, progresso, i * INTERVALO_ENTRE_LOGINS),
            name=f"N{i + 1}",
            daemon=True,
        )
        for i in range(quantidade)
    ]
    for t in threads:
        t.start()

    try:
        while any(t.is_alive() for t in threads):
            time.sleep(1)
    except KeyboardInterrupt:
        print("\nInterrompido: cada navegador termina o paciente atual e fecha "
              "(Ctrl+C de novo para sair na hora).")
        PARAR.set()
        for t in threads:
            t.join()
    finally:
        parar_relogio.set()

    if MOTIVO_PARADA:
        print("\n" + MOTIVO_PARADA[0])
    elif fila.empty():
        print("\nCadastro de orçamentos concluído!")
    imprimir_resumo(progresso, fila)


if __name__ == "__main__":
    executar()
