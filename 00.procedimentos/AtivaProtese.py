import openpyxl
import os
import re
import time
import traceback
import unicodedata
from datetime import datetime
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.action_chains import ActionChains


# Marca a opção "Prótese" (controla-protese) em todo procedimento com Protese = "Sim"
ARQUIVO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Template Procedimentos, Grupos e Produtos.xlsx")
ABA = "Procedimentos"
COLUNA_STATUS = "Status Protese"
CLINICA = "BRASILIA - GRUPO MANDIC"

URL_BASE = "https://app.clinicanasnuvens.com.br"
URL_LISTA = f"{URL_BASE}/procedimentos"
PASTA_ERROS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "erros")

# Quantos procedimentos processar nesta execução (None = todos). Use 1 para testar.
LIMITE = None

# True = roda sem abrir a janela do Chrome
HEADLESS = True

# Segundos que cada espera aguarda a página/botão aparecer
ESPERA = 30


def _salvar_planilha(wb):
    # Salva num temporário e troca pelo original (os.replace é atômico):
    # uma interrupção no meio nunca deixa o .xlsx corrompido.
    tmp = ARQUIVO + ".tmp"
    wb.save(tmp)
    os.replace(tmp, ARQUIVO)


def _idx_coluna(cabecalho, nome, obrigatoria=True):
    # Busca a coluna pelo nome, ignorando maiúsculas/minúsculas e espaços nas pontas.
    alvo = nome.strip().lower()
    for i, c in enumerate(cabecalho):
        if c and str(c).strip().lower() == alvo:
            return i + 1
    if obrigatoria:
        raise ValueError(f"Coluna '{nome}' não encontrada na aba '{ABA}'. Colunas disponíveis: {cabecalho}")
    return None


def _texto(valor):
    """Célula -> texto: 2.0 (número no Excel) -> '2'; vazio -> ''."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def normalizar(valor):
    """Minúsculas, sem acentos e com espaços simples — para comparar nomes."""
    valor = unicodedata.normalize("NFKD", _texto(valor))
    valor = "".join(c for c in valor if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", valor).strip().lower()


def carregar_procedimentos():
    """Devolve (col_status, {nome: [linhas]}) dos procedimentos com Protese = Sim ainda não processados."""
    wb = openpyxl.load_workbook(ARQUIVO)
    ws = wb[ABA]
    cabecalho = [cell.value for cell in ws[1]]

    col_procedimento = _idx_coluna(cabecalho, "Nome do procedimento")
    col_protese      = _idx_coluna(cabecalho, "Protese")

    col_status = _idx_coluna(cabecalho, COLUNA_STATUS, obrigatoria=False)
    if col_status is None:
        col_status = len(cabecalho) + 1
        ws.cell(row=1, column=col_status, value=COLUNA_STATUS)
        _salvar_planilha(wb)

    procedimentos = {}
    for row in ws.iter_rows(min_row=2):
        num_row = row[0].row
        procedimento = _texto(ws.cell(row=num_row, column=col_procedimento).value)
        protese      = normalizar(ws.cell(row=num_row, column=col_protese).value)
        status       = normalizar(ws.cell(row=num_row, column=col_status).value)

        if not procedimento or protese != "sim" or status in ("ativado", "erro"):
            continue
        procedimentos.setdefault(procedimento, []).append(num_row)

    return col_status, procedimentos


def marcar_status(linhas, col_status, valor):
    wb = openpyxl.load_workbook(ARQUIVO)
    ws = wb[ABA]
    for linha in linhas:
        ws.cell(row=linha, column=col_status, value=valor)
    _salvar_planilha(wb)


def salvar_print_erro(driver, nome):
    try:
        os.makedirs(PASTA_ERROS, exist_ok=True)
        nome = re.sub(r"[^\w\-]+", "_", nome)[:60]
        base = os.path.join(PASTA_ERROS, f"protese_{nome}_{datetime.now():%Y%m%d_%H%M%S}")
        driver.save_screenshot(base + ".png")
        with open(base + ".html", "w", encoding="utf-8") as f:
            f.write(driver.page_source)
        print(f"  Print salvo em {base}.png")
    except Exception:
        pass


def clicar(driver, elemento):
    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", elemento)
    driver.execute_script("arguments[0].click();", elemento)


# ── Passos na plataforma ───────────────────────────────────────────────────────

def abrir_procedimento(driver, nome_proc):
    """Pesquisa o procedimento e abre o Editar da linha com o nome exato (não só o 1º resultado)."""
    driver.get(URL_LISTA)
    campo = WebDriverWait(driver, ESPERA).until(EC.element_to_be_clickable((By.ID, "nomeTipoProcedimento")))
    campo.clear()
    campo.send_keys(nome_proc)

    botao = next((b for b in driver.find_elements(By.CSS_SELECTOR, "button[type='submit'].btn-light-primary")
                  if "pesquisar" in b.text.lower()), None)
    if botao is None:
        raise Exception("botão 'Pesquisar' não encontrado")
    clicar(driver, botao)

    # Espera a lista recarregar com o resultado da pesquisa
    WebDriverWait(driver, ESPERA).until(lambda d: "nomeTipoProcedimento=" in d.current_url
                                        or d.find_elements(By.CSS_SELECTOR, "a.visualizar-procedimento"))
    time.sleep(1.5)

    links = driver.find_elements(By.CSS_SELECTOR, "a.visualizar-procedimento")
    if not links:
        raise Exception("procedimento não encontrado na pesquisa")

    # A pesquisa é por "contém": 'Barra 3X3' também traz 'Barra 3X3 Reta'. Fica com o nome exato.
    alvo = normalizar(nome_proc)
    for link in links:
        try:
            linha = link.find_element(By.XPATH, "./ancestor::tr[1]")
        except Exception:
            continue
        if any(normalizar(celula.text) == alvo for celula in linha.find_elements(By.CSS_SELECTOR, "td")):
            clicar(driver, link)
            return

    if len(links) == 1:
        print("  Aviso: nome na tela não bateu exatamente; usando o único resultado da pesquisa.")
        clicar(driver, links[0])
        return
    raise Exception(f"{len(links)} resultados na pesquisa e nenhum com o nome exato")


def ativar_protese(driver):
    """Marca 'controla-protese' e salva. Devolve False se já estava marcado (nada a salvar)."""
    checkbox = WebDriverWait(driver, ESPERA).until(EC.presence_of_element_located((By.ID, "controla-protese")))
    time.sleep(1)
    if checkbox.is_selected():
        return False

    clicar(driver, checkbox)
    time.sleep(0.5)
    if not checkbox.is_selected():
        raise Exception("não foi possível marcar a opção Prótese")

    botao_salvar = next(
        (b for b in driver.find_elements(By.CSS_SELECTOR, "button[type='submit'][name='_method'][value='post']")
         if b.is_displayed() and "salvar" in b.text.lower()),
        None,
    )
    if botao_salvar is None:
        raise Exception("botão 'Salvar' não encontrado")

    url_procedimento = driver.current_url
    clicar(driver, botao_salvar)
    time.sleep(2)

    # Confere: reabre o procedimento e vê se a opção ficou salva
    driver.get(url_procedimento)
    checkbox = WebDriverWait(driver, ESPERA).until(EC.presence_of_element_located((By.ID, "controla-protese")))
    if not checkbox.is_selected():
        raise Exception("salvou, mas a opção Prótese não ficou marcada (o sistema recusou o Salvar?)")
    return True


def executar():
    col_status, procedimentos = carregar_procedimentos()

    if not procedimentos:
        print("Nenhum procedimento pendente com Protese = Sim.")
        return
    print(f"{len(procedimentos)} procedimento(s) com Protese = Sim pendente(s).")

    opcoes = webdriver.ChromeOptions()
    if HEADLESS:
        opcoes.add_argument("--headless=new")
    opcoes.add_argument("--window-size=1920,1080")
    driver = webdriver.Chrome(options=opcoes)
    wait = WebDriverWait(driver, ESPERA)

    try:
        # ── Login ──────────────────────────────────────────────────────────────
        driver.get(f"{URL_BASE}/")

        campo_email = wait.until(EC.presence_of_element_located(
            (By.XPATH, "/html/body/div[1]/div/div[2]/div[1]/div/div[3]/form/div[1]/input")
        ))
        campo_email.send_keys("mmerlo++++dpegoraro@bionexo.com")
        driver.find_element(By.XPATH, "/html/body/div[1]/div/div[2]/div[1]/div/div[3]/form/div[2]/button").click()

        campo_senha = wait.until(EC.element_to_be_clickable((By.ID, "password")))
        campo_senha.click()
        campo_senha.send_keys("CNN@foguete")
        driver.find_element(By.XPATH, "/html/body/div[1]/div/div[2]/div[1]/div/div[4]/form/div[4]/button").click()
        time.sleep(5)

        # ── Seleciona a clínica ────────────────────────────────────────────────
        for popup in driver.find_elements(By.XPATH, "//button[text()='X']"):
            if popup.is_displayed():
                driver.execute_script("arguments[0].click();", popup)
                time.sleep(1)
                break

        botao = wait.until(EC.presence_of_element_located((By.ID, "clinica-selecionada")))
        ActionChains(driver).move_to_element(botao).click(botao).perform()
        time.sleep(1)

        campo_busca = wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//input[@placeholder='Pesquisar clínica']")
        ))
        campo_busca.click()
        campo_busca.send_keys(CLINICA)
        time.sleep(2)

        opcao = wait.until(EC.presence_of_element_located(
            (By.XPATH, f"//a[contains(@class,'trocar-clinica') and contains(.,'{CLINICA}')]")
        ))
        driver.execute_script("arguments[0].click();", opcao)
        time.sleep(5)

        # ── Loop principal ─────────────────────────────────────────────────────
        total = len(procedimentos)
        inicio = time.time()
        ativados = ja_estavam = erros = 0

        for idx, (nome_proc, linhas) in enumerate(procedimentos.items(), start=1):
            if LIMITE is not None and idx > LIMITE:
                break
            print(f"\n[{idx}/{total}] {nome_proc}")

            try:
                abrir_procedimento(driver, nome_proc)
                if ativar_protese(driver):
                    ativados += 1
                    print("  Prótese ativada e salva.")
                else:
                    ja_estavam += 1
                    print("  Prótese já estava ativada — nada a salvar.")
                marcar_status(linhas, col_status, "ativado")

            except Exception as e:
                erros += 1
                print(f"  ERRO: {type(e).__name__}: {e}")
                salvar_print_erro(driver, nome_proc)
                marcar_status(linhas, col_status, "erro")

            media = (time.time() - inicio) / idx
            print(f"  Média {media:.1f}s/procedimento | restantes: {total - idx}")

        print(f"\nFim! {ativados} ativado(s), {ja_estavam} já estavam ativados, {erros} erro(s). "
              f"Tempo total: {time.time() - inicio:.0f}s")

    except Exception:
        salvar_print_erro(driver, "sessao")
        raise

    finally:
        driver.quit()


if __name__ == "__main__":
    # Reinicia em caso de erro de sessão (linhas já marcadas são puladas)
    while True:
        try:
            executar()
            break
        except Exception as e:
            print(f"\nERRO: {type(e).__name__}: {e}")
            traceback.print_exc(limit=3)
            print("Reiniciando em 10 segundos...")
            time.sleep(10)
