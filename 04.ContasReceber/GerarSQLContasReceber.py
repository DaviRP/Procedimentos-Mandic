"""
Gerador de script SQL para faturamento em lote (sp_criar_faturamento_completo).

Lê os títulos em aberto e monta, para cada linha, uma chamada
`CALL sp_criar_faturamento_completo(...)`, concatenando tudo em um único
arquivo .sql pronto para ser executado no banco.

Entrada (ARQUIVO_ENTRADA):
  - .xlsx no formato do Modelo_ContasReceber.xlsx (recomendado): títulos na aba
    "Titulos" e os valores fixos na aba "Configuração" (substituem os da config
    abaixo quando preenchidos).
  - .csv com as colunas PG, Valor, Data, nome, IDLAN (usa a config abaixo).

Mapeamento de colunas -> parâmetros da procedure:
    p_tenant_id        -> Configuração: Tenant ID
    p_valor            -> coluna "Valor"
    p_data_faturamento -> coluna "Data" (DD/MM/AAAA -> AAAA-MM-DD)
    p_codcliente       -> subselect usando a coluna "PG" como numerocontrole
    p_codconta         -> Configuração: Conta (codconta)
    p_numerodocumento  -> Configuração: Número do documento
    p_tipodocumento    -> Configuração: Tipo de documento
    p_descricao        -> coluna "nome"
    p_codcategoria     -> Configuração: Categoria (codcategoria)
    p_idtotvs          -> coluna "IDLAN"
    p_idcoligadatotvs  -> Configuração: Coligada TOTVS
"""

import csv
from datetime import date, datetime
from pathlib import Path

# =============================================================================
# CONFIGURAÇÃO — ajustar antes de rodar
# =============================================================================

ARQUIVO_ENTRADA = "Modelo_ContasReceber.xlsx"  # .xlsx do modelo ou .csv
ARQUIVO_SAIDA = "ContasReceber_Insert.sql"
ENCODING_CSV = "latin-1"  # só para .csv (o arquivo original não é UTF-8)

# Usados para .csv, ou quando a aba Configuração do .xlsx estiver vazia
TENANT_ID = 14570
CODCONTA = 25746
NUMERODOCUMENTO = "31232"
CODCATEGORIA = 616956
IDCOLIGADATOTVS = "1"
TIPODOCUMENTO = "CARTAO_DEBITO"

# =============================================================================

COLUNAS = ("PG", "Valor", "Data", "nome", "IDLAN")

# Aba Configuração do modelo: célula -> chave da config
CELULAS_CONFIG = {
    "B5": "TENANT_ID",
    "B6": "CODCONTA",
    "B7": "NUMERODOCUMENTO",
    "B8": "CODCATEGORIA",
    "B9": "IDCOLIGADATOTVS",
    "B10": "TIPODOCUMENTO",
}


def escapar_sql(valor: str) -> str:
    """Escapa aspas simples para uso seguro dentro de literais SQL."""
    return str(valor).replace("'", "''")


def converter_data(data_str: str) -> str:
    """Converte DD/MM/AAAA -> AAAA-MM-DD."""
    return datetime.strptime(data_str.strip(), "%d/%m/%Y").strftime("%Y-%m-%d")


def converter_valor(valor_str: str) -> str:
    """'115,60' / '115.6' -> '115.60' (o SQL precisa de ponto como decimal)."""
    return f"{float(valor_str.strip().replace(',', '.')):.2f}"


def _texto_celula(valor) -> str:
    """Valor de célula do Excel -> texto no mesmo formato que viria de um CSV."""
    if valor is None:
        return ""
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def config_padrao() -> dict:
    return {
        "TENANT_ID": TENANT_ID,
        "CODCONTA": CODCONTA,
        "NUMERODOCUMENTO": NUMERODOCUMENTO,
        "CODCATEGORIA": CODCATEGORIA,
        "IDCOLIGADATOTVS": IDCOLIGADATOTVS,
        "TIPODOCUMENTO": TIPODOCUMENTO,
    }


def ler_csv(caminho: Path):
    with caminho.open(newline="", encoding=ENCODING_CSV) as f:
        return list(csv.DictReader(f)), config_padrao()


def ler_xlsx(caminho: Path):
    """Lê o Modelo_ContasReceber.xlsx: títulos da aba Titulos + config da aba Configuração."""
    import openpyxl

    wb = openpyxl.load_workbook(caminho, data_only=True)

    config = config_padrao()
    if "Configuração" in wb.sheetnames:
        aba_config = wb["Configuração"]
        for celula, chave in CELULAS_CONFIG.items():
            valor = _texto_celula(aba_config[celula].value)
            if valor:
                config[chave] = valor

    aba = wb["Titulos"]
    cabecalho = [_texto_celula(c.value) for c in aba[1]]
    faltando = [c for c in COLUNAS if c not in cabecalho]
    if faltando:
        raise SystemExit(f"Colunas não encontradas na aba Titulos: {faltando} (não renomeie o cabeçalho do modelo)")
    indices = {c: cabecalho.index(c) for c in COLUNAS}

    linhas = []
    for valores in aba.iter_rows(min_row=2, values_only=True):
        linha = {c: _texto_celula(valores[i]) for c, i in indices.items()}
        if not any(linha.values()):
            continue  # linha vazia do modelo
        linhas.append(linha)
    wb.close()
    return linhas, config


def montar_call(linha: dict, config: dict) -> str:
    pg = escapar_sql(linha["PG"].strip())
    valor = converter_valor(linha["Valor"])
    data_faturamento = converter_data(linha["Data"])
    descricao = escapar_sql(linha["nome"].strip())
    idtotvs = int(linha["IDLAN"].strip())

    return f"""CALL sp_criar_faturamento_completo(
    {config['TENANT_ID']},
    {valor},
    '{data_faturamento}',
    (
        SELECT cf.codpessoa
        FROM CLIENTE_FORNECEDOR cf
        WHERE cf.codpessoa_fk = (
            SELECT p.codpessoa_fk
            FROM PACIENTE p
            WHERE p.numerocontrole = '{pg}'
              AND p.tenant_id = {config['TENANT_ID']}
        )
    ),
    {config['CODCONTA']},
    '{escapar_sql(config['NUMERODOCUMENTO'])}',
    '{escapar_sql(config['TIPODOCUMENTO'])}',
    '{descricao}',
    {config['CODCATEGORIA']},
    {idtotvs},
    '{escapar_sql(config['IDCOLIGADATOTVS'])}',
    @pai_id,
    @fat_id
);"""


def main():
    pasta_script = Path(__file__).resolve().parent
    caminho_entrada = pasta_script / ARQUIVO_ENTRADA
    caminho_saida = pasta_script / ARQUIVO_SAIDA

    if caminho_entrada.suffix.lower() == ".xlsx":
        linhas, config = ler_xlsx(caminho_entrada)
    else:
        linhas, config = ler_csv(caminho_entrada)

    # Campos numéricos da config precisam ser números (vão sem aspas no SQL)
    for chave in ("TENANT_ID", "CODCONTA", "CODCATEGORIA"):
        try:
            config[chave] = int(str(config[chave]).strip())
        except ValueError:
            raise SystemExit(f"Configuração inválida: {chave} = {config[chave]!r} (precisa ser número inteiro)")

    print("Configuração usada:")
    for chave, valor in config.items():
        print(f"  {chave}: {valor}")

    blocos = []
    ignoradas = []
    idlans_vistos = {}

    for i, linha in enumerate(linhas, start=2):  # linha 1 = cabeçalho
        pg = (linha.get("PG") or "").strip()
        valor = (linha.get("Valor") or "").strip()
        data = (linha.get("Data") or "").strip()
        nome = (linha.get("nome") or "").strip()
        idlan = (linha.get("IDLAN") or "").strip()

        if not pg or not valor or not data or not nome or not idlan:
            ignoradas.append((i, "campo obrigatório vazio (PG/Valor/Data/nome/IDLAN)"))
            continue

        try:
            converter_data(data)
        except ValueError:
            ignoradas.append((i, f"data inválida: {data!r}"))
            continue

        try:
            if float(converter_valor(valor)) <= 0:
                ignoradas.append((i, f"valor deve ser maior que zero: {valor!r}"))
                continue
        except ValueError:
            ignoradas.append((i, f"valor inválido: {valor!r}"))
            continue

        if len(nome) > 255:
            ignoradas.append((i, f"nome com {len(nome)} caracteres (máximo 255)"))
            continue

        if not idlan.isdigit():
            ignoradas.append((i, f"IDLAN deve ser número inteiro: {idlan!r}"))
            continue

        if idlan in idlans_vistos:
            ignoradas.append((i, f"IDLAN {idlan} repetido (já usado na linha {idlans_vistos[idlan]})"))
            continue
        idlans_vistos[idlan] = i

        comentario = f"-- Linha {i} | PG {pg} | {nome}"
        blocos.append(f"{comentario}\n{montar_call(linha, config)}")

    cabecalho = (
        "-- Script gerado automaticamente por GerarSQLContasReceber.py\n"
        f"-- Gerado em: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
        f"-- Entrada: {caminho_entrada.name} | Tenant: {config['TENANT_ID']}\n"
        f"-- Total de chamadas: {len(blocos)} | Ignoradas: {len(ignoradas)}\n\n"
    )

    conteudo = cabecalho + "\n\n".join(blocos) + "\n"
    caminho_saida.write_text(conteudo, encoding="utf-8")

    print(f"\nLinhas lidas: {len(linhas)}")
    print(f"CALLs geradas: {len(blocos)}")
    print(f"Arquivo gerado: {caminho_saida.resolve()}")

    if ignoradas:
        print(f"\nLinhas ignoradas ({len(ignoradas)}):")
        for numero, motivo in ignoradas:
            print(f"  - linha {numero}: {motivo}")


if __name__ == "__main__":
    main()
