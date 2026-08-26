import html
import re
import time
import unicodedata
import openpyxl
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.support.ui import WebDriverWait, Select
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import WebDriverException


ARQUIVO = "SOLICITACOES_SERVICOS_v2.xlsx"
CLINICA = "FORTALEZA - GRUPO MANDIC"

# ── Filtro de teste ────────────────────────────────────────────────────────
# Enquanto True, só processa os pacientes listados em NOMES_TESTE.
# Depois de validar, é só voltar para False para rodar a planilha inteira.
FILTRAR_TESTE = True
NOMES_TESTE = [
    "Francisca Nogueira Vieira",
    "MARIA GORETTI CARDOSO DE SOUZA",
    "Luely de Souza Santos",
    "ELIONE FERNANDES DA SILVA",
    "ANA EDITE DE ASSIS CARVALHO",
    "MARIA ANTONIA DOS SANTOS GOMES",
    "MARCIA NEIDE DE OLIVEIRA",
    "LIDUINA SOARES DA SILVA",
    "JÚLIA CASTRO VASCONCELOS",
    "CINEIDE CRISTINA DE LIMA FONTINELLE",
    "MARILISA RUGGERI QUEIROZ",
    "LÚCIO LAERTTI RIOS CAJAZEIRAS",
]


def _normalizar_nome(nome):
    # NFC unifica acentos que às vezes vêm em forma decomposta (ex: "U" + acento
    # combinante em vez do caractere "Ú" já composto), o que faria strings
    # visualmente iguais não baterem na comparação.
    nome = unicodedata.normalize("NFC", nome or "")
    nome = re.sub(r"\s+", " ", nome)
    return nome.strip().upper()


_NOMES_TESTE_NORMALIZADOS = {_normalizar_nome(n) for n in NOMES_TESTE}


def corrigir_encoding(texto):
    """Corrige mojibake: bytes UTF-8 que foram decodificados como cp1252/Latin-1
    (ex: 'OrÃ§amento' -> 'Orçamento', 'PRÃ“TESE' -> 'PRÓTESE').

    A fonte original usa cp1252 (Windows-1252) — por isso aspas curvas como em
    'Ã“' aparecem no lugar de acentos (0x93 em cp1252 é “, que não existe em
    Latin-1 puro). Tenta cp1252 primeiro e cai para Latin-1 como reserva.

    Só tenta corrigir se houver sinal de mojibake (o 'Ã'/'Â' é a marca
    característica dessa dupla-codificação); sem esse sinal, o texto já está
    correto e é devolvido sem alteração. Se algum trecho tiver bytes realmente
    inválidos (registro corrompido na origem), corrige o resto do texto e
    substitui só o trecho quebrado, em vez de descartar a correção inteira."""
    if not texto:
        return ""
    if "Ã" not in texto and "Â" not in texto:
        return texto

    for codificacao in ("cp1252", "latin1"):
        try:
            return texto.encode(codificacao, errors="strict").decode("utf-8", errors="strict")
        except (UnicodeDecodeError, UnicodeEncodeError):
            continue

    try:
        return texto.encode("cp1252", errors="replace").decode("utf-8", errors="replace")
    except UnicodeEncodeError:
        return texto


def tratar_texto(texto_html):
    """Converte o HTML bruto do campo 'texto' em texto plano legível, multi-linha:
    corrige o encoding, transforma <br/> em quebra de linha, remove as demais tags
    e desfaz entidades HTML (&amp;, &nbsp;, ...)."""
    texto = corrigir_encoding(texto_html or "")

    texto = re.sub(r"<br\s*/?>", "\n", texto, flags=re.IGNORECASE)
    texto = re.sub(r"<[^>]+>", "", texto)

    texto = html.unescape(texto)
    texto = texto.replace("\xa0", " ")

    linhas = [linha.strip() for linha in texto.split("\n")]
    linhas = [linha for linha in linhas if linha]
    return "\n".join(linhas)


def carregar_solicitacoes():
    """Lê SOLICITACOES_SERVICOS_v2.xlsx e devolve uma lista de dicts, uma entrada
    por linha da planilha (cada linha vira uma observação a ser cadastrada)."""
    wb = openpyxl.load_workbook(ARQUIVO, data_only=True)
    ws = wb.active

    cabecalho = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    idx = {nome: i for i, nome in enumerate(cabecalho)}

    obrigatorias = ["codigoImportacao", "data", "codigoPaciente", "PG", "nome", "cpf", "texto"]
    faltando = [c for c in obrigatorias if c not in idx]
    if faltando:
        raise ValueError(f"Coluna(s) obrigatória(s) não encontrada(s) na planilha: {faltando}")

    solicitacoes = []
    for linha in ws.iter_rows(min_row=2, values_only=True):
        pg = str(linha[idx["PG"]] or "").strip()
        if not pg:
            continue

        nome = str(linha[idx["nome"]] or "").strip()

        if FILTRAR_TESTE:
            if _normalizar_nome(nome) not in _NOMES_TESTE_NORMALIZADOS:
                continue

        solicitacoes.append({
            "codigo_importacao": str(linha[idx["codigoImportacao"]] or "").strip(),
            "data": linha[idx["data"]],
            "codigo_paciente": str(linha[idx["codigoPaciente"]] or "").strip(),
            "pg": pg,
            "nome": nome,
            "cpf": str(linha[idx["cpf"]] or "").strip(),
            "texto": tratar_texto(linha[idx["texto"]]),
        })

    return solicitacoes


def executar():
    solicitacoes = carregar_solicitacoes()

    if not solicitacoes:
        print("Nenhuma solicitação encontrada.")
        return

    if FILTRAR_TESTE:
        print(f"[FILTRO DE TESTE ATIVO] {len(solicitacoes)} solicitação(ões) selecionada(s) "
              f"de {len(NOMES_TESTE)} nome(s) em NOMES_TESTE.")
        for sol in solicitacoes:
            print(f"  - {sol['nome']} (PG {sol['pg']})")

    opcoes = webdriver.ChromeOptions()
    opcoes.add_argument("--window-size=1920,1080")
    opcoes.add_argument("--window-position=-32000,-32000")
    driver = webdriver.Chrome(options=opcoes)
    wait = WebDriverWait(driver, 15)

    try:
        # ── Login ──────────────────────────────────────────────────────────────
        driver.get("https://app.clinicanasnuvens.com.br/")

        campo_email = wait.until(EC.presence_of_element_located(
            (By.XPATH, "/html/body/div[1]/div/div[2]/div[1]/div/div[3]/form/div[1]/input")
        ))
        campo_email.send_keys("dpegoraro++++dpegoraro@bionexo.com")
        driver.find_element(By.XPATH, "/html/body/div[1]/div/div[2]/div[1]/div/div[3]/form/div[2]/button").click()

        campo_senha = wait.until(EC.element_to_be_clickable((By.ID, "password")))
        campo_senha.click()
        campo_senha.send_keys("Bi0n3xdpegrr06")
        driver.find_element(By.XPATH, "/html/body/div[1]/div/div[2]/div[1]/div/div[4]/form/div[4]/button").click()
        time.sleep(5)

        # ── Seleciona a clínica ────────────────────────────────────────────────
        for popup in driver.find_elements(By.XPATH, "//button[text()='X']"):
            if popup.is_displayed():
                driver.execute_script("arguments[0].click();", popup)
                time.sleep(1)
                break

        botao = wait.until(EC.presence_of_element_located((By.ID, "clinica-selecionada")))
        ActionChains(driver).move_to_element(botao).click(botao).perform()
        time.sleep(1)

        campo_busca = wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//input[@placeholder='Pesquisar clínica']")
        ))
        campo_busca.click()
        campo_busca.send_keys(CLINICA)
        time.sleep(2)

        opcao = wait.until(EC.presence_of_element_located(
            (By.XPATH, f"//a[contains(@class,'trocar-clinica') and contains(.,'{CLINICA}')]")
        ))
        driver.execute_script("arguments[0].click();", opcao)
        time.sleep(5)

        # ── Loop principal ─────────────────────────────────────────────────────
        total = len(solicitacoes)

        for idx, sol in enumerate(solicitacoes, start=1):
            nome_paciente = sol["nome"]
            print(f"\n[{idx}/{total}] PG {sol['pg']} — {nome_paciente}")

            try:
                # 1. Vai para a lista de pacientes e abre o filtro
                driver.get("https://app.clinicanasnuvens.com.br/pacientes")
                time.sleep(2)

                botao_filtrar = wait.until(EC.element_to_be_clickable((By.ID, "btn-drawer-filtrosearch-paciente")))
                driver.execute_script("arguments[0].click();", botao_filtrar)
                time.sleep(1)

                # 2. Busca pelo PG (número de controle)
                campo_pg = wait.until(EC.element_to_be_clickable((By.ID, "pesquisa-numero-controle")))
                campo_pg.clear()
                campo_pg.send_keys(sol["pg"])

                botao_visualizar = wait.until(EC.element_to_be_clickable((By.ID, "btn-visualizar-pacientes")))
                driver.execute_script("arguments[0].click();", botao_visualizar)
                time.sleep(2)

                # 3. Abre o perfil do primeiro paciente da lista
                link_paciente = wait.until(EC.element_to_be_clickable(
                    (By.CSS_SELECTOR, "a[aria-label='Visualizar perfil do paciente']")
                ))
                driver.execute_script("arguments[0].click();", link_paciente)
                time.sleep(2)

                print("  Paciente localizado.")

                # 4. Abre a aba Observações
                botao_observacoes = wait.until(EC.element_to_be_clickable(
                    (By.CSS_SELECTOR, "button.nav-link.observacoes.historico")
                ))
                driver.execute_script("arguments[0].click();", botao_observacoes)
                time.sleep(1)

                # 5. Abre o modal "Adicionar observação"
                # Pode haver mais de um elemento com esse ID na página (templates
                # duplicados/AJAX); pega o primeiro que estiver de fato visível.
                botao_adicionar = wait.until(lambda d: d.execute_script("""
                    var els = document.querySelectorAll("[id='adicionar-observacao']");
                    for (var i = 0; i < els.length; i++) {
                        var el = els[i];
                        var r = el.getBoundingClientRect();
                        if (r.width > 0 && r.height > 0 && el.offsetParent !== null) {
                            return el;
                        }
                    }
                    return null;
                """))
                botao_adicionar.click()
                time.sleep(1)

                # 6. Digita o texto tratado no editor (CKEditor), linha a linha,
                # simulando digitação real para não quebrar o modelo interno do editor
                editor = wait.until(EC.element_to_be_clickable(
                    (By.CSS_SELECTOR, "div.ck-editor__editable[contenteditable='true']")
                ))
                editor.click()

                linhas = sol["texto"].split("\n")
                for i, linha in enumerate(linhas):
                    editor.send_keys(linha)
                    if i < len(linhas) - 1:
                        editor.send_keys(Keys.ENTER)
                time.sleep(0.5)

                # 7. Seleciona o status "Fechado"
                status_select = wait.until(EC.presence_of_element_located((By.ID, "pacienteObservacaoStatus")))
                Select(status_select).select_by_value("FECHADO")

                # 8. Salva
                botao_salvar = wait.until(EC.element_to_be_clickable(
                    (By.XPATH, "//div[contains(@class,'modal') and contains(@class,'show')]"
                               "//button[@type='submit' and contains(.,'Salvar')]")
                ))
                driver.execute_script("arguments[0].click();", botao_salvar)
                time.sleep(1.5)

                print("  Observação salva.")

            except WebDriverException as e:
                mensagem = str(e).lower()
                sessao_morta = any(s in mensagem for s in (
                    "no such window", "invalid session id", "target window already closed",
                    "chrome not reachable", "disconnected",
                ))
                if sessao_morta:
                    print(f"  ERRO FATAL: sessão do navegador caiu (PG {sol['pg']}): {e}")
                    print("  Reiniciando o navegador do zero...")
                    raise

                print(f"  ERRO ao processar PG {sol['pg']} ({nome_paciente}): {e}")
                try:
                    driver.get("https://app.clinicanasnuvens.com.br/pacientes")
                    time.sleep(2)
                    print("  De volta à lista de pacientes — indo para o próximo.")
                except Exception:
                    pass
                continue

            except Exception as e:
                print(f"  ERRO ao processar PG {sol['pg']} ({nome_paciente}): {e}")
                try:
                    driver.get("https://app.clinicanasnuvens.com.br/pacientes")
                    time.sleep(2)
                    print("  De volta à lista de pacientes — indo para o próximo.")
                except Exception:
                    pass
                continue

        print("\nImportação de observações concluída!")

    except Exception:
        raise

    finally:
        driver.quit()


if __name__ == "__main__":
    # Reinicia tudo do zero em caso de qualquer erro
    while True:
        try:
            executar()
            break
        except Exception as e:
            print(f"\nERRO: {e}")
            print("Reiniciando em 10 segundos...")
            time.sleep(10)
